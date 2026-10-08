from datetime import date, timedelta
import csv
import io
import mimetypes
import os
import re
import secrets
from urllib.parse import urlsplit

from django.contrib import messages
from django.contrib.auth.models import User
from django.contrib.auth.decorators import login_required
from django.conf import settings
from django.core.exceptions import PermissionDenied
from django.core.files.images import get_image_dimensions
from django.db import transaction
from django.db.models import OuterRef, Q, Subquery, Value
from django.db.models.functions import Replace
from django.http import FileResponse, Http404, HttpResponse, HttpResponseBadRequest, JsonResponse, StreamingHttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_POST
from django.views.decorators.csrf import ensure_csrf_cookie
from django.views.decorators.debug import sensitive_post_parameters
from docx import Document

from pool_service.client_crm_models import ClientContact
from pool_service.client_queries import active_clients
from pool_service.client_phone_matching import clients_by_phones
from pool_service.phone_utils import canonical_phone_value, format_phone, normalize_phone
from pool_service.communication_avito import (
    AvitoError,
    authorized_account_id as avito_authorized_account_id,
    subscribe_webhook as avito_subscribe_webhook,
    sync_recent_messages as avito_sync_recent_messages,
    unsubscribe_webhook as avito_unsubscribe_webhook,
    verify_messenger_access as avito_verify_messenger_access,
    webhook_subscriptions as avito_webhook_subscriptions,
)
from pool_service.communication_forms import CommunicationConnectionForm, TelephonyConnectionForm
from pool_service.communication_models import (
    AvitoCredential, ChannelConnection, CommunicationChannel, Conversation,
    ConversationAssignment, ConversationMessage, ConversationReadState,
    CallAnalysis, MessageAttachment, PhoneCall, TelephonyConnection,
)
from pool_service.communication_recordings import looks_like_audio_file
from pool_service.communication_secrets import encrypt_secret
from pool_service.communication_services import conversation_capability, optimize_message_image, organization_access
from pool_service.services.call_ai import PROCESSING_STALE_MINUTES, request_call_analysis, start_requested_call_analysis_worker
from pool_service.models import Client, Notification, OrganizationAccess



def _context(request, capability):
    access = organization_access(request.user)
    if not access or not conversation_capability(request.user, capability, access.organization):
        raise PermissionDenied
    return access.organization


def _can_access_manual_recordings(user, organization):
    access = organization_access(user, organization)
    return bool(
        access
        and (access.role in {"owner", "admin"} or user.is_superuser)
    )


def _owner_communications_context(request):
    access = organization_access(request.user)
    if not access or not (
        access.role in {"owner", "admin"} or request.user.is_superuser
    ):
        raise PermissionDenied
    return access.organization


def _is_manual_recording_call(call):
    return call.source_kind == PhoneCall.SOURCE_UPLOADED


def _authorize_call_access(request, call):
    access = organization_access(request.user, call.organization)
    if not access:
        raise PermissionDenied
    if _is_manual_recording_call(call):
        if not (
            access.role in {"owner", "admin"}
            or request.user.is_superuser
        ):
            raise PermissionDenied
        return access.organization
    if not conversation_capability(request.user, "can_listen_calls", access.organization):
        raise PermissionDenied
    if (
        not conversation_capability(request.user, "can_view_all_calls", access.organization)
        and call.employee_id != request.user.id
        and call.peer_employee_id != request.user.id
    ):
        raise PermissionDenied
    return access.organization


def _wake_uploaded_audio_worker_if_needed(organization, *, include_fresh_pending=False):
    now = timezone.now()
    stale_before = now - timedelta(minutes=PROCESSING_STALE_MINUTES)
    pending_before = now if include_fresh_pending else now - timedelta(seconds=30)
    needs_worker = (
        CallAnalysis.objects.filter(
            call__organization=organization,
            call__source_kind=PhoneCall.SOURCE_UPLOADED,
            requested_at__isnull=False,
        )
        .filter(
            Q(
                status=CallAnalysis.STATUS_PENDING,
                requested_at__lte=pending_before,
            )
            | Q(
                status=CallAnalysis.STATUS_PROCESSING,
                processing_started_at__lt=stale_before,
            )
        )
        .exists()
    )
    if needs_worker:
        start_requested_call_analysis_worker()
    return needs_worker


@login_required
def conversations(request):
    organization = _context(request, "can_view_conversations")
    latest_message = ConversationMessage.objects.filter(conversation=OuterRef("pk")).order_by("-created_at", "-pk")
    items = list(
        Conversation.objects.filter(organization=organization)
        .select_related("connection__channel", "assignee", "website_request")
        .annotate(last_message_preview=Subquery(latest_message.values("body")[:1]))
    )
    states = {state.conversation_id: state.last_read_at for state in ConversationReadState.objects.filter(user=request.user, conversation__organization=organization)}
    for item in items:
        incoming = item.messages.filter(direction=ConversationMessage.DIRECTION_IN)
        last_read = states.get(item.pk)
        item.unread_count = incoming.filter(created_at__gt=last_read).count() if last_read else incoming.count()
    selected = None
    selected_id = request.GET.get("conversation")
    if selected_id:
        selected = next((item for item in items if str(item.uuid) == selected_id), None)
        if selected is None:
            raise Http404
    elif items:
        selected = items[0]
    if selected:
        state, _ = ConversationReadState.objects.get_or_create(conversation=selected, user=request.user)
        state.last_read_at = timezone.now()
        state.save(update_fields=["last_read_at"])
        Notification.objects.filter(
            user=request.user,
            organization=selected.organization,
            kind="communication",
            action_url__contains=str(selected.uuid),
            is_resolved=False,
        ).update(is_read=True, is_resolved=True, resolved_at=timezone.now())
    return render(request, "pool_service/communications/conversations.html", {
        "active_tab": "communications", "conversations": items, "selected": selected,
        "can_reply": conversation_capability(request.user, "can_reply_conversations", organization),
        "can_take": conversation_capability(request.user, "can_take_conversation", organization),
        "can_assign": conversation_capability(request.user, "can_assign_conversation", organization),
        "can_access_manual_recordings": _can_access_manual_recordings(request.user, organization),
        "staff": OrganizationAccess.objects.filter(organization=organization).exclude(role="accountant").select_related("user"),
    })


@login_required
@require_POST
@transaction.atomic
def conversation_reply(request, conversation_uuid):
    organization = _context(request, "can_reply_conversations")
    conversation = get_object_or_404(
        Conversation.objects.select_related("connection__channel"),
        uuid=conversation_uuid,
        organization=organization,
    )
    channel = get_object_or_404(
        CommunicationChannel.objects.select_for_update(),
        pk=conversation.connection.channel_id,
        organization=organization,
    )
    connection = get_object_or_404(
        ChannelConnection.objects.select_for_update(),
        pk=conversation.connection_id,
        channel=channel,
    )
    if not connection.is_active or not channel.is_active:
        messages.error(request, "Канал или подключение отключено. Отправка ответа недоступна.")
        return redirect(f"/communications/?conversation={conversation.uuid}")
    body = request.POST.get("body", "").strip()
    files = request.FILES.getlist("attachments")
    if len(body) > 10000:
        messages.error(request, "Сообщение не должно превышать 10 000 символов.")
        return redirect(f"/communications/?conversation={conversation.uuid}")
    if files and conversation.connection.channel.kind == CommunicationChannel.KIND_AVITO:
        messages.error(request, "Отправка фото и файлов в Авито пока не подключена.")
        return redirect(f"/communications/?conversation={conversation.uuid}")
    if len(files) > 10 or sum(upload.size for upload in files) > 25 * 1024 * 1024:
        messages.error(request, "Можно приложить до 10 файлов общим размером не более 25 МБ.")
        return redirect(f"/communications/?conversation={conversation.uuid}")
    if any(upload.size > 10 * 1024 * 1024 for upload in files):
        messages.error(request, "Размер одного файла не должен превышать 10 МБ.")
        return redirect(f"/communications/?conversation={conversation.uuid}")
    if not body and not files:
        messages.error(request, "Введите сообщение или добавьте файл.")
        return redirect("communications_conversations")
    for uploaded in files:
        content_type = uploaded.content_type or "application/octet-stream"
        if content_type.startswith("image/"):
            try:
                dimensions = get_image_dimensions(uploaded)
                if not dimensions or not all(dimensions):
                    raise ValueError("invalid dimensions")
                width, height = dimensions
                if width > 12000 or height > 12000 or width * height > 25_000_000:
                    messages.error(request, f"Изображение {uploaded.name} имеет слишком большое разрешение.")
                    return redirect(f"/communications/?conversation={conversation.uuid}")
                uploaded.seek(0)
            except Exception:
                messages.error(request, f"Файл {uploaded.name} не является корректным изображением.")
                return redirect(f"/communications/?conversation={conversation.uuid}")
    message = ConversationMessage.objects.create(
        conversation=conversation,
        direction="out",
        body=body,
        sender_name=request.user.get_full_name() or request.user.username,
        sent_by=request.user,
        delivery_status=ConversationMessage.DELIVERY_PENDING,
    )
    for uploaded in files:
        content_type = uploaded.content_type or "application/octet-stream"
        attachment = MessageAttachment.objects.create(message=message, original=uploaded, original_name=uploaded.name[:255], content_type=content_type, original_size=uploaded.size)
        try:
            optimize_message_image(attachment)
        except Exception:
            # The original remains available even when a safe preview cannot be produced.
            pass
    conversation.last_message_at = message.created_at
    conversation.save(update_fields=["last_message_at", "updated_at"])
    # Channel adapters send queued outbound messages; the persisted message is the source of truth.
    return redirect(f"/communications/?conversation={conversation.uuid}")


@login_required
@require_POST
@transaction.atomic
def conversation_take(request, conversation_uuid):
    organization = _context(request, "can_take_conversation")
    conversation = get_object_or_404(Conversation.objects.select_for_update(), uuid=conversation_uuid, organization=organization)
    if conversation.assignee_id and conversation.assignee_id != request.user.id:
        raise PermissionDenied
    conversation.assignee = request.user
    conversation.status = Conversation.STATUS_ACTIVE
    conversation.save(update_fields=["assignee", "status", "updated_at"])
    ConversationAssignment.objects.create(conversation=conversation, assignee=request.user, changed_by=request.user)
    Notification.objects.filter(
        organization=conversation.organization,
        kind="communication",
        action_url__contains=str(conversation.uuid),
        is_resolved=False,
    ).exclude(user=request.user).update(is_resolved=True, resolved_at=timezone.now())
    return redirect(f"/communications/?conversation={conversation.uuid}")


@login_required
@require_POST
@transaction.atomic
def conversation_update(request, conversation_uuid):
    organization = _context(request, "can_view_conversations")
    conversation = get_object_or_404(Conversation.objects.select_for_update(), uuid=conversation_uuid, organization=organization)
    status = request.POST.get("status")
    assignee_id = request.POST.get("assignee")
    if assignee_id is not None:
        if not conversation_capability(request.user, "can_assign_conversation", organization):
            raise PermissionDenied
        assignee = None
        if assignee_id:
            assignee = get_object_or_404(User, pk=assignee_id, organizationaccess__organization=organization)
        conversation.assignee = assignee
        conversation.save(update_fields=["assignee", "updated_at"])
        ConversationAssignment.objects.create(conversation=conversation, assignee=assignee, changed_by=request.user)
    if status in dict(Conversation.STATUS_CHOICES):
        if not conversation_capability(request.user, "can_take_conversation", organization):
            raise PermissionDenied
        if conversation.assignee_id not in (None, request.user.id) and not conversation_capability(request.user, "can_assign_conversation", organization):
            raise PermissionDenied
        conversation.status = status
        conversation.save(update_fields=["status", "updated_at"])
    return redirect(f"/communications/?conversation={conversation.uuid}")


@login_required
def attachment_download(request, attachment_id):
    organization = _context(request, "can_view_conversations")
    attachment = get_object_or_404(MessageAttachment, pk=attachment_id, message__conversation__organization=organization)
    response = FileResponse(
        attachment.original.open("rb"),
        as_attachment=True,
        filename=attachment.original_name,
        content_type="application/octet-stream",
    )
    response["X-Content-Type-Options"] = "nosniff"
    return response


def _call_filter_dates(request):
    try:
        date_from = (
            date.fromisoformat(request.GET["date_from"])
            if request.GET.get("date_from")
            else None
        )
        date_to = (
            date.fromisoformat(request.GET["date_to"])
            if request.GET.get("date_to")
            else None
        )
    except ValueError as exc:
        raise ValueError("Некорректный период.") from exc
    if date_from and date_to and date_from > date_to:
        raise ValueError("Начало периода не может быть позже окончания.")
    return date_from, date_to


def _phone_digits_expression():
    expression = "phone_number"
    for separator in ("+", " ", "-", "(", ")", "."):
        expression = Replace(
            expression,
            Value(separator),
            Value(""),
        )
    return expression


def _apply_date_filters(queryset, request):
    date_from, date_to = _call_filter_dates(request)
    if date_from:
        queryset = queryset.filter(started_at__date__gte=date_from)
    if date_to:
        queryset = queryset.filter(started_at__date__lte=date_to)
    return queryset


def _filter_telephony_calls(request, queryset, organization, can_view_all):
    queryset = _apply_date_filters(queryset, request)

    employee_id = (request.GET.get("employee") or "").strip()
    if employee_id and can_view_all:
        if (
            not employee_id.isdigit()
            or not OrganizationAccess.objects.filter(
                organization=organization,
                user_id=employee_id,
            ).exists()
        ):
            raise ValueError("Некорректный сотрудник.")
        queryset = queryset.filter(
            Q(employee_id=employee_id) | Q(peer_employee_id=employee_id)
        )

    client_id = (request.GET.get("client") or "").strip()
    if client_id:
        if not client_id.isdigit():
            raise ValueError("Некорректный клиент.")
        client = active_clients(
            Client.objects.filter(
                organization=organization,
                pk=int(client_id),
            )
        ).first()
        if client is None:
            raise ValueError("Некорректный клиент.")

        phone_values = set(
            ClientContact.objects.filter(
                client=client,
                kind=ClientContact.KIND_PHONE,
            )
            .exclude(match_value="")
            .values_list("match_value", flat=True)
        )
        legacy_phone = normalize_phone(client.phone)
        if legacy_phone:
            phone_values.add(legacy_phone)

        client_filter = Q(client_id=client.pk)
        if phone_values:
            queryset = queryset.annotate(
                phone_number_digits=_phone_digits_expression()
            )
            for phone_value in phone_values:
                digits = "".join(ch for ch in phone_value if ch.isdigit())
                if digits:
                    client_filter |= Q(
                        phone_number_digits__icontains=digits[-10:]
                    )
        queryset = queryset.filter(client_filter)

    direction = request.GET.get("direction")
    if direction in ("in", "out", PhoneCall.DIRECTION_INTERNAL):
        queryset = queryset.filter(direction=direction)
    if request.GET.get("missed"):
        queryset = queryset.filter(result=PhoneCall.RESULT_MISSED)

    query = (request.GET.get("q") or "").strip()
    if query:
        normalized_query = normalize_phone(query)
        canonical_query = canonical_phone_value(query)
        raw_phone_digits = "".join(ch for ch in query if ch.isdigit())
        normalized_digits = "".join(ch for ch in normalized_query if ch.isdigit())
        phone_query = (
            normalized_digits[-10:]
            if len(normalized_digits) >= 10
            else raw_phone_digits[-10:]
        )
        queryset = queryset.annotate(
            phone_number_digits=_phone_digits_expression()
        )
        phone_filter = (
            Q(phone_number__icontains=query)
            | Q(phone_number__icontains=canonical_query)
            | Q(contact_name__icontains=query)
            | Q(client__name__icontains=query)
            | Q(provider_user__icontains=query)
            | Q(provider_extension__icontains=query)
            | Q(peer_provider_user__icontains=query)
            | Q(peer_provider_extension__icontains=query)
            | Q(employee_profile__display_name__icontains=query)
            | Q(peer_employee_profile__display_name__icontains=query)
            | Q(employee__first_name__icontains=query)
            | Q(employee__last_name__icontains=query)
            | Q(peer_employee__first_name__icontains=query)
            | Q(peer_employee__last_name__icontains=query)
        )
        if phone_query:
            phone_filter |= Q(phone_number_digits__icontains=phone_query)
        queryset = queryset.filter(phone_filter)

    return queryset


def _filter_uploaded_audio(request, queryset, organization):
    queryset = _apply_date_filters(queryset, request)

    client_id = (request.GET.get("client") or "").strip()
    if client_id:
        if not client_id.isdigit():
            raise ValueError("Некорректный клиент.")
        if not Client.objects.filter(
            pk=int(client_id),
            organization=organization,
        ).exists():
            raise ValueError("Некорректный клиент.")
        queryset = queryset.filter(client_id=int(client_id))

    query = (request.GET.get("q") or "").strip()
    if query:
        normalized_query = normalize_phone(query)
        canonical_query = canonical_phone_value(query)
        phone_digits = "".join(ch for ch in normalized_query if ch.isdigit())
        queryset = queryset.annotate(
            phone_number_digits=_phone_digits_expression()
        )
        search_filter = (
            Q(client__name__icontains=query)
            | Q(contact_name__icontains=query)
            | Q(phone_number__icontains=query)
            | Q(phone_number__icontains=canonical_query)
            | Q(recording_file__icontains=query)
        )
        if phone_digits:
            search_filter |= Q(
                phone_number_digits__icontains=phone_digits[-10:]
            )
        queryset = queryset.filter(search_filter)

    return queryset


def _resolve_export_clients(calls, organization):
    unresolved = [
        call.phone_number
        for call in calls
        if not call.client_id and normalize_phone(call.phone_number)
    ]
    phone_matches = clients_by_phones(organization, unresolved)
    resolved = {}
    for call in calls:
        if call.client_id:
            resolved[call.pk] = call.client
            continue
        matches = phone_matches.get(normalize_phone(call.phone_number), [])
        resolved[call.pk] = matches[0] if len(matches) == 1 else None
    return resolved


def _transcript_export_records(queryset, organization):
    calls = list(
        queryset.filter(
            analysis__status=CallAnalysis.STATUS_READY,
        )
        .exclude(analysis__transcript="")
        .select_related(
            "analysis",
            "client",
            "employee",
            "employee_profile",
            "peer_employee",
            "peer_employee_profile",
        )
        .order_by("started_at", "pk")
    )
    resolved_clients = _resolve_export_clients(calls, organization)
    records = []
    for call in calls:
        client = resolved_clients.get(call.pk)
        if call.employee_profile:
            employee = call.employee_profile.display_name
        elif call.employee:
            employee = call.employee.get_full_name() or call.employee.username
        else:
            employee = call.provider_user or call.provider_extension

        if call.direction == PhoneCall.DIRECTION_INTERNAL:
            if call.peer_employee_profile:
                peer_employee = call.peer_employee_profile.display_name
            elif call.peer_employee:
                peer_employee = (
                    call.peer_employee.get_full_name()
                    or call.peer_employee.username
                )
            else:
                peer_employee = (
                    call.contact_name
                    or call.peer_provider_user
                    or call.peer_provider_extension
                )
            employee = " → ".join(
                value for value in (employee, peer_employee) if value
            )

        started_at = call.started_at
        if timezone.is_aware(started_at):
            started_at = timezone.localtime(started_at)

        records.append({
            "date_time": started_at.strftime("%d.%m.%Y %H:%M"),
            "client": (
                ""
                if call.direction == PhoneCall.DIRECTION_INTERNAL
                else (client.name if client else (call.contact_name or ""))
            ),
            "phone": (
                ""
                if call.direction == PhoneCall.DIRECTION_INTERNAL
                else format_phone(call.phone_number)
            ),
            "employee": employee,
            "direction": call.get_direction_display(),
            "duration": call.duration_seconds,
            "file_name": call.recording_filename,
            "summary": (call.analysis.summary or "").strip(),
            "transcript": (call.analysis.transcript or "").strip(),
        })
    return records


def _safe_csv_cell(value):
    text = "" if value is None else str(value)
    if text[:1] in {"=", "+", "-", "@"}:
        return "'" + text
    return text


def _transcript_export_response(records, export_format, source_kind):
    label = "аудиофайлов" if source_kind == PhoneCall.SOURCE_UPLOADED else "звонков"
    today = timezone.localdate().isoformat()
    base_name = f"transcripts_{source_kind}_{today}"

    if export_format == "txt":
        parts = [f"Расшифровки {label}", f"Количество: {len(records)}", ""]
        for index, record in enumerate(records, start=1):
            parts.extend([
                f"{index}. {record['date_time']}",
                f"Клиент: {record['client'] or '—'}",
                f"Телефон: {record['phone'] or '—'}",
            ])
            if record["employee"]:
                parts.append(f"Сотрудник: {record['employee']}")
            if source_kind == PhoneCall.SOURCE_TELEPHONY:
                parts.append(f"Направление: {record['direction']}")
                parts.append(f"Длительность: {record['duration']} сек.")
            elif record["file_name"]:
                parts.append(f"Файл: {record['file_name']}")
            parts.extend([
                "",
                "Итог разговора:",
                record["summary"] or "—",
                "",
                "Полная расшифровка:",
                record["transcript"] or "—",
                "",
                "=" * 80,
                "",
            ])
        response = HttpResponse(
            "\ufeff" + "\n".join(parts),
            content_type="text/plain; charset=utf-8",
        )
        response["Content-Disposition"] = f'attachment; filename="{base_name}.txt"'
        return response

    if export_format == "csv":
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer)
        writer.writerow([
            "Дата и время",
            "Клиент",
            "Телефон",
            "Сотрудник",
            "Направление",
            "Длительность, сек.",
            "Файл",
            "Итог разговора",
            "Полная расшифровка",
        ])
        for record in records:
            writer.writerow([
                _safe_csv_cell(record["date_time"]),
                _safe_csv_cell(record["client"]),
                _safe_csv_cell(record["phone"]),
                _safe_csv_cell(record["employee"]),
                _safe_csv_cell(
                    record["direction"]
                    if source_kind == PhoneCall.SOURCE_TELEPHONY
                    else ""
                ),
                record["duration"] if source_kind == PhoneCall.SOURCE_TELEPHONY else "",
                _safe_csv_cell(record["file_name"]),
                _safe_csv_cell(record["summary"]),
                _safe_csv_cell(record["transcript"]),
            ])
        response = HttpResponse(
            "\ufeff" + buffer.getvalue(),
            content_type="text/csv; charset=utf-8",
        )
        response["Content-Disposition"] = f'attachment; filename="{base_name}.csv"'
        return response

    if export_format == "docx":
        document = Document()
        document.add_heading(f"Расшифровки {label}", 0)
        document.add_paragraph(f"Количество: {len(records)}")

        if not records:
            document.add_paragraph("По выбранным фильтрам готовых расшифровок нет.")

        for index, record in enumerate(records, start=1):
            title_parts = [record["date_time"]]
            if record["client"]:
                title_parts.append(record["client"])
            elif record["phone"]:
                title_parts.append(record["phone"])
            document.add_heading(
                f"{index}. " + " — ".join(title_parts),
                level=1,
            )

            metadata = []
            if record["phone"]:
                metadata.append(("Телефон", record["phone"]))
            if record["employee"]:
                metadata.append(("Сотрудник", record["employee"]))
            if source_kind == PhoneCall.SOURCE_TELEPHONY:
                metadata.append(("Направление", record["direction"]))
                metadata.append(("Длительность", f"{record['duration']} сек."))
            elif record["file_name"]:
                metadata.append(("Файл", record["file_name"]))

            if metadata:
                table = document.add_table(rows=0, cols=2)
                table.style = "Table Grid"
                for key, value in metadata:
                    cells = table.add_row().cells
                    cells[0].text = key
                    cells[1].text = str(value)

            document.add_heading("Итог разговора", level=2)
            document.add_paragraph(record["summary"] or "—")
            document.add_heading("Полная расшифровка", level=2)
            document.add_paragraph(record["transcript"] or "—")

        output = io.BytesIO()
        document.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )
        response["Content-Disposition"] = f'attachment; filename="{base_name}.docx"'
        return response

    return HttpResponseBadRequest("Неизвестный формат выгрузки.")


@login_required
def call_transcripts_export(request, source_kind):
    export_format = (request.GET.get("format") or "txt").lower()
    if export_format not in {"txt", "csv", "docx"}:
        return HttpResponseBadRequest("Неизвестный формат выгрузки.")

    if source_kind == PhoneCall.SOURCE_UPLOADED:
        organization = _owner_communications_context(request)
        queryset = PhoneCall.objects.filter(
            organization=organization,
            source_kind=PhoneCall.SOURCE_UPLOADED,
        )
        try:
            queryset = _filter_uploaded_audio(request, queryset, organization)
        except ValueError as exc:
            return HttpResponseBadRequest(str(exc))
    elif source_kind == PhoneCall.SOURCE_TELEPHONY:
        access = organization_access(request.user)
        if not access:
            raise PermissionDenied
        organization = access.organization
        if not conversation_capability(
            request.user,
            "can_listen_calls",
            organization,
        ):
            raise PermissionDenied
        can_view_all = conversation_capability(
            request.user,
            "can_view_all_calls",
            organization,
        )
        queryset = PhoneCall.objects.filter(
            organization=organization,
            source_kind=PhoneCall.SOURCE_TELEPHONY,
        )
        if not can_view_all:
            queryset = queryset.filter(
                Q(employee=request.user) | Q(peer_employee=request.user)
            )
        try:
            queryset = _filter_telephony_calls(
                request,
                queryset,
                organization,
                can_view_all,
            )
        except ValueError as exc:
            return HttpResponseBadRequest(str(exc))
    else:
        return HttpResponseBadRequest("Некорректный источник.")

    records = _transcript_export_records(queryset, organization)
    return _transcript_export_response(records, export_format, source_kind)


@login_required
def calls(request):
    access = organization_access(request.user)
    if not access or not (
        conversation_capability(request.user, "can_view_own_calls", access.organization)
        or conversation_capability(request.user, "can_view_all_calls", access.organization)
    ):
        raise PermissionDenied
    organization = access.organization
    can_view_all = conversation_capability(request.user, "can_view_all_calls", organization)
    queryset = PhoneCall.objects.filter(
        organization=organization,
        source_kind=PhoneCall.SOURCE_TELEPHONY,
    ).select_related(
        "employee",
        "employee_profile",
        "peer_employee",
        "peer_employee_profile",
        "client",
        "analysis",
    ).defer("analysis__transcript")
    if not can_view_all:
        queryset = queryset.filter(
            Q(employee=request.user) | Q(peer_employee=request.user)
        )
    try:
        queryset = _filter_telephony_calls(
            request,
            queryset,
            organization,
            can_view_all,
        )
    except ValueError as exc:
        return HttpResponseBadRequest(str(exc))
    employees = OrganizationAccess.objects.filter(
        organization=organization
    ).select_related("user")
    clients = active_clients(
        Client.objects.filter(organization=organization)
    ).order_by("name", "id")
    calls = list(queryset[:500])
    unresolved_phone_values = [
        call.phone_number
        for call in calls
        if not call.client_id and normalize_phone(call.phone_number)
    ]
    phone_matches = clients_by_phones(organization, unresolved_phone_values)
    for call in calls:
        call.phone_display = format_phone(call.phone_number)
        if call.client_id:
            call.resolved_client = call.client
            call.ambiguous_clients = []
            continue
        matches = phone_matches.get(normalize_phone(call.phone_number), [])
        if len(matches) == 1:
            call.resolved_client = matches[0]
            call.ambiguous_clients = []
        else:
            call.resolved_client = None
            call.ambiguous_clients = matches
    return render(request, "pool_service/communications/calls.html", {
        "active_tab": "communications",
        "calls": calls,
        "employees": employees,
        "clients": clients,
        "can_listen": conversation_capability(request.user, "can_listen_calls", organization),
        "can_view_all": can_view_all,
        "can_access_manual_recordings": _can_access_manual_recordings(request.user, organization),
        "manual_archive": False,
    })


@login_required
def manual_recordings(request):
    organization = _owner_communications_context(request)
    _wake_uploaded_audio_worker_if_needed(
        organization,
        include_fresh_pending=True,
    )
    queryset = PhoneCall.objects.filter(
        organization=organization,
        source_kind=PhoneCall.SOURCE_UPLOADED,
    ).select_related(
        "client",
        "analysis",
    ).defer("analysis__transcript")

    try:
        queryset = _filter_uploaded_audio(request, queryset, organization)
    except ValueError as exc:
        return HttpResponseBadRequest(str(exc))

    clients = active_clients(
        Client.objects.filter(organization=organization)
    ).order_by("name", "id")
    return render(request, "pool_service/communications/calls.html", {
        "active_tab": "communications",
        "calls": queryset[:500],
        "clients": clients,
        "can_listen": True,
        "can_view_all": True,
        "can_access_manual_recordings": True,
        "manual_archive": True,
    })


MANUAL_CALL_ALLOWED_EXTENSIONS = {".mp3", ".mp4", ".mpeg", ".mpga", ".m4a", ".wav", ".webm"}
MANUAL_CALL_MAX_FILES = 50


@login_required
@require_POST
def call_recording_upload(request):
    organization = _owner_communications_context(request)

    files = request.FILES.getlist("recordings")
    if not files:
        messages.error(request, "Выберите хотя бы один аудиофайл.")
        return redirect("communication_manual_recordings")
    if len(files) > MANUAL_CALL_MAX_FILES:
        messages.error(request, f"За один раз можно загрузить не более {MANUAL_CALL_MAX_FILES} записей.")
        return redirect("communication_manual_recordings")

    client = None
    client_id = (request.POST.get("client") or "").strip()
    if client_id:
        if not client_id.isdigit():
            return HttpResponseBadRequest("Некорректный клиент.")
        client = get_object_or_404(Client, pk=client_id, organization=organization)

    max_bytes = int(getattr(settings, "COMMUNICATION_RECORDING_MAX_BYTES", 50 * 1024 * 1024))
    rejected = []
    for uploaded in files:
        extension = os.path.splitext(uploaded.name or "")[1].lower()
        if extension not in MANUAL_CALL_ALLOWED_EXTENSIONS:
            rejected.append(f"{uploaded.name}: неподдерживаемый формат")
            continue
        if uploaded.size <= 0:
            rejected.append(f"{uploaded.name}: пустой файл")
            continue
        if uploaded.size > max_bytes:
            rejected.append(f"{uploaded.name}: файл больше допустимого размера")
            continue
        prefix = uploaded.read(32)
        uploaded.seek(0)
        if not looks_like_audio_file(prefix, extension):
            rejected.append(f"{uploaded.name}: файл не похож на поддерживаемую аудиозапись")
    if rejected:
        messages.error(request, "Не загружено: " + "; ".join(rejected[:5]))
        return redirect("communication_manual_recordings")

    created = 0
    for uploaded in files:
        call = PhoneCall(
            organization=organization,
            source_kind=PhoneCall.SOURCE_UPLOADED,
            connection=None,
            external_id=f"upload-{secrets.token_hex(16)}",
            client=client,
            contact_name=(client.name if client else ""),
            phone_number=(canonical_phone_value(client.phone) if client else ""),
            direction=PhoneCall.DIRECTION_IN,
            started_at=timezone.now(),
            duration_seconds=0,
            result=PhoneCall.RESULT_ANSWERED,
            recording_status=PhoneCall.RECORDING_STORED,
            recording_downloaded_at=timezone.now(),
        )
        call.recording_file.save(uploaded.name, uploaded, save=False)
        call.save()
        created += 1

    messages.success(
        request,
        f"Загружено записей: {created}. Расшифровка не запускалась.",
    )
    return redirect("communication_manual_recordings")


@login_required
@require_POST
def call_analysis_bulk(request):
    organization = _owner_communications_context(request)
    raw_ids = request.POST.getlist("call_ids")
    call_ids = [int(value) for value in raw_ids if value.isdigit()]
    if not call_ids:
        messages.info(request, "Выберите записи для расшифровки.")
        return redirect("communication_manual_recordings")

    queryset = PhoneCall.objects.filter(
        pk__in=call_ids,
        organization=organization,
        source_kind=PhoneCall.SOURCE_UPLOADED,
    ).exclude(recording_file="").filter(
        Q(analysis__isnull=True)
        | Q(analysis__status=CallAnalysis.STATUS_FAILED)
        | Q(
            analysis__status=CallAnalysis.STATUS_PENDING,
            analysis__requested_at__isnull=True,
        )
    )
    queued = 0
    for call_id in queryset.order_by("pk").values_list("pk", flat=True):
        if request_call_analysis(call_id, allow_reanalysis=False):
            queued += 1

    if queued:
        if start_requested_call_analysis_worker():
            messages.success(request, f"Поставлено на расшифровку: {queued}.")
        else:
            messages.warning(
                request,
                f"Поставлено в очередь: {queued}. Резервный обработчик подхватит записи позже.",
            )
    else:
        messages.info(request, "Новые записи в очередь не добавлены.")
    return redirect("communication_manual_recordings")


@login_required
def call_analysis_status(request, call_id):
    access = organization_access(request.user)
    if not access:
        raise PermissionDenied
    call = get_object_or_404(
        PhoneCall.objects.select_related("analysis", "connection"),
        pk=call_id,
        organization=access.organization,
    )
    _authorize_call_access(request, call)
    analysis = getattr(call, "analysis", None)
    if analysis is None:
        return JsonResponse({"status": "none"})
    if _is_manual_recording_call(call):
        _wake_uploaded_audio_worker_if_needed(call.organization)
    return JsonResponse({
        "status": analysis.status,
        "error": analysis.error,
    })


@login_required
def call_analysis_transcript(request, call_id):
    access = organization_access(request.user)
    if not access:
        raise PermissionDenied
    call = get_object_or_404(
        PhoneCall.objects.select_related("analysis", "connection"),
        pk=call_id,
        organization=access.organization,
    )
    _authorize_call_access(request, call)
    analysis = getattr(call, "analysis", None)
    if not analysis or analysis.status != CallAnalysis.STATUS_READY:
        raise Http404
    return JsonResponse({"transcript": analysis.transcript})


@login_required
@require_POST
def call_analysis_retry(request, call_id):
    access = organization_access(request.user)
    if not access:
        raise PermissionDenied
    call = get_object_or_404(
        PhoneCall.objects.select_related("connection"),
        pk=call_id,
        organization=access.organization,
    )
    _authorize_call_access(request, call)
    redirect_name = (
        "communication_manual_recordings"
        if _is_manual_recording_call(call)
        else "communications_calls"
    )
    if not call.recording_file:
        messages.error(request, "Сначала должна быть сохранена запись звонка.")
        return redirect(redirect_name)
    if request_call_analysis(call.pk):
        if start_requested_call_analysis_worker():
            messages.success(request, "Расшифровка и анализ запущены.")
        else:
            messages.warning(
                request,
                "Звонок поставлен в очередь. Резервный обработчик подхватит его позже.",
            )
    else:
        messages.info(request, "Этот звонок уже обрабатывается.")
    return redirect(redirect_name)


def _recording_range_iterator(file_handle, start, length, chunk_size=64 * 1024):
    try:
        file_handle.seek(start)
        remaining = length
        while remaining > 0:
            chunk = file_handle.read(min(chunk_size, remaining))
            if not chunk:
                break
            remaining -= len(chunk)
            yield chunk
    finally:
        file_handle.close()


@login_required
def call_recording(request, call_id):
    access = organization_access(request.user)
    if not access:
        raise PermissionDenied
    call = get_object_or_404(
        PhoneCall.objects.select_related("connection"),
        pk=call_id,
        organization=access.organization,
    )
    _authorize_call_access(request, call)
    if not call.recording_file:
        raise Http404

    try:
        file_path = call.recording_file.path
        file_size = os.path.getsize(file_path)
    except (OSError, ValueError):
        raise Http404

    download_requested = request.GET.get("download") == "1"
    range_header = request.headers.get("Range", "").strip()

    if download_requested or not range_header:
        content_type = mimetypes.guess_type(file_path)[0] or "application/octet-stream"
        response = FileResponse(
            open(file_path, "rb"),
            content_type=content_type,
            as_attachment=download_requested,
            filename=os.path.basename(file_path),
        )
        response["Accept-Ranges"] = "bytes"
        response["Content-Length"] = str(file_size)
        response["Cache-Control"] = "private, max-age=3600"
        response["X-Content-Type-Options"] = "nosniff"
        return response

    match = re.fullmatch(r"bytes=(\d*)-(\d*)", range_header)
    if not match:
        response = HttpResponse(status=416)
        response["Content-Range"] = f"bytes */{file_size}"
        return response

    start_raw, end_raw = match.groups()
    if not start_raw and not end_raw:
        response = HttpResponse(status=416)
        response["Content-Range"] = f"bytes */{file_size}"
        return response

    if start_raw:
        start = int(start_raw)
        end = int(end_raw) if end_raw else file_size - 1
    else:
        suffix_length = int(end_raw)
        if suffix_length <= 0:
            response = HttpResponse(status=416)
            response["Content-Range"] = f"bytes */{file_size}"
            return response
        start = max(file_size - suffix_length, 0)
        end = file_size - 1

    if start >= file_size or start > end:
        response = HttpResponse(status=416)
        response["Content-Range"] = f"bytes */{file_size}"
        return response

    end = min(end, file_size - 1)
    length = end - start + 1
    response = StreamingHttpResponse(
        _recording_range_iterator(open(file_path, "rb"), start, length),
        status=206,
        content_type=mimetypes.guess_type(file_path)[0] or "application/octet-stream",
    )
    response["Content-Length"] = str(length)
    response["Content-Range"] = f"bytes {start}-{end}/{file_size}"
    response["Accept-Ranges"] = "bytes"
    response["Cache-Control"] = "private, max-age=3600"
    response["X-Content-Type-Options"] = "nosniff"
    return response


def _connection_setup_result(request, connection, secret_value, *, secret_kind):
    if secret_kind == "website":
        endpoint = request.build_absolute_uri(
            reverse("website_chat_message", args=[connection.public_id])
        )
        title = "Подключение сайта создано"
        secret_label = "Bearer-токен"
        instructions = (
            "Скопируйте токен сейчас и сохраните в защищённой серверной "
            "конфигурации сайта. Service2 больше его не покажет."
        )
        callback_url = ""
    else:
        endpoint = ""
        title = "Подключение Авито настроено"
        secret_label = "Webhook-токен"
        instructions = (
            "Скопируйте webhook URL сейчас и зарегистрируйте его в Авито. "
            "Токен хранится в Service2 только в виде хеша и повторно не показывается."
        )
        callback_url = request.build_absolute_uri(
            reverse("avito_webhook", args=[connection.public_id, secret_value])
        )
    response = render(
        request,
        "pool_service/communications/connection_secret.html",
        {
            "active_tab": "communications",
            "title": title,
            "connection": connection,
            "secret_label": secret_label,
            "secret_value": secret_value,
            "endpoint": endpoint,
            "callback_url": callback_url,
            "instructions": instructions,
        },
    )
    response["Cache-Control"] = "no-store"
    response["Pragma"] = "no-cache"
    response["Referrer-Policy"] = "no-referrer"
    return response


def _telephony_setup_result(request, telephony, provider_connection, token):
    callback_url = request.build_absolute_uri(
        reverse("megafon_webhook", args=[provider_connection.public_id])
    )
    response = render(
        request,
        "pool_service/communications/telephony_secret.html",
        {
            "active_tab": "communications",
            "telephony": telephony,
            "callback_url": callback_url,
            "crm_token": token,
        },
    )
    response["Cache-Control"] = "no-store"
    response["Pragma"] = "no-cache"
    response["Referrer-Policy"] = "no-referrer"
    return response


def _telephony_provider_connection(telephony):
    channel = _communication_provider_channel(
        telephony.organization, CommunicationChannel.KIND_MEGAFON
    )
    connection, _ = ChannelConnection.objects.get_or_create(
        channel=channel,
        external_id=telephony.external_id,
        defaults={
            "name": telephony.name,
            "is_active": telephony.is_active,
        },
    )
    changed = []
    if connection.name != telephony.name:
        connection.name = telephony.name
        changed.append("name")
    if connection.is_active != telephony.is_active:
        connection.is_active = telephony.is_active
        changed.append("is_active")
    if changed:
        connection.save(update_fields=changed)
    return connection



def _save_megafon_api_settings(provider_connection, form):
    settings_data = dict(provider_connection.settings or {})
    settings_data["megafon_api_base_url"] = form.cleaned_data["ats_base_url"]
    if form.cleaned_data.get("ats_api_key"):
        settings_data["megafon_api_key_encrypted"] = encrypt_secret(
            form.cleaned_data["ats_api_key"]
        )
    provider_connection.settings = settings_data
    provider_connection.save(update_fields=["settings"])


def _communication_provider_channel(organization, kind):
    channel = (
        CommunicationChannel.objects.filter(organization=organization, kind=kind)
        .order_by("pk")
        .first()
    )
    if channel:
        return channel
    return CommunicationChannel.objects.create(
        organization=organization,
        kind=kind,
        name=dict(CommunicationChannel.KIND_CHOICES)[kind],
        is_active=True,
    )


@login_required
def channels(request):
    organization = _context(request, "can_manage_channels")
    communication_channels = list(
        CommunicationChannel.objects.filter(organization=organization)
        .prefetch_related("connections")
        .order_by("kind", "pk")
    )
    avito_credential_ids = set(
        AvitoCredential.objects.filter(
            connection__channel__organization=organization
        ).values_list("connection_id", flat=True)
    )
    provider_connections = {"website": [], "avito": []}
    provider_channels = {"website": [], "avito": []}
    for channel in communication_channels:
        if channel.kind not in provider_connections:
            continue
        provider_channels[channel.kind].append(channel)
        for connection in channel.connections.all():
            connection.configuration_ready = bool(connection.api_token_hash)
            if channel.kind == CommunicationChannel.KIND_AVITO:
                connection.configuration_ready = connection.pk in avito_credential_ids
                connection.avito_webhook_status = (connection.settings or {}).get(
                    "avito_webhook_status", "not_connected"
                )
                connection.avito_webhook_checked_at = (connection.settings or {}).get(
                    "avito_webhook_checked_at", ""
                )
                connection.avito_webhook_last_received_at = (connection.settings or {}).get(
                    "avito_webhook_last_received_at", ""
                )
                connection.avito_webhook_last_result = (connection.settings or {}).get(
                    "avito_webhook_last_result", ""
                )
                connection.avito_webhook_error = (connection.settings or {}).get(
                    "avito_webhook_error", ""
                )
                connection.avito_webhook_last_error = (connection.settings or {}).get(
                    "avito_webhook_last_error", ""
                )
                connection.avito_pull_last_checked_at = (connection.settings or {}).get(
                    "avito_pull_last_checked_at", ""
                )
                connection.avito_pull_last_created = (connection.settings or {}).get(
                    "avito_pull_last_created", ""
                )
                connection.avito_pull_last_error = (connection.settings or {}).get(
                    "avito_pull_last_error", ""
                )
            provider_connections[channel.kind].append(connection)

    providers = [
        {
            "kind": "website",
            "title": "Сайт",
            "description": "Чат и заявки с сайта через защищённый API Service2.",
            "channels": provider_channels["website"],
            "connections": provider_connections["website"],
        },
        {
            "kind": "avito",
            "title": "Авито",
            "description": "Входящие сообщения и ответы менеджеров через Avito API.",
            "channels": provider_channels["avito"],
            "connections": provider_connections["avito"],
        },
    ]
    telephony_connections = list(
        TelephonyConnection.objects.filter(organization=organization)
        .exclude(external_id="manual-upload")
        .order_by("pk")
    )
    megafon_connections = {
        item.external_id: item
        for item in ChannelConnection.objects.filter(
            channel__organization=organization,
            channel__kind=CommunicationChannel.KIND_MEGAFON,
        )
    }
    for telephony in telephony_connections:
        provider_connection = megafon_connections.get(telephony.external_id)
        telephony.provider_connection = provider_connection
        telephony.megafon_webhook_configured = bool(
            provider_connection and provider_connection.api_token_hash
        )
        telephony.megafon_api_configured = bool(
            provider_connection
            and (provider_connection.settings or {}).get("megafon_api_base_url")
            and (provider_connection.settings or {}).get("megafon_api_key_encrypted")
        )
        telephony.megafon_last_received_at = (
            (provider_connection.settings or {}).get("megafon_last_received_at", "")
            if provider_connection
            else ""
        )
        telephony.megafon_last_result = (
            (provider_connection.settings or {}).get("megafon_last_result", "")
            if provider_connection
            else ""
        )
        telephony.megafon_last_event_at = (
            (provider_connection.settings or {}).get("megafon_last_event_at", "")
            if provider_connection
            else ""
        )
        telephony.megafon_last_event_type = (
            (provider_connection.settings or {}).get("megafon_last_event_type", "")
            if provider_connection
            else ""
        )
        telephony.megafon_last_event_phone = (
            (provider_connection.settings or {}).get("megafon_last_event_phone", "")
            if provider_connection
            else ""
        )
        telephony.megafon_last_event_user = (
            (provider_connection.settings or {}).get("megafon_last_event_user", "")
            if provider_connection
            else ""
        )
        telephony.megafon_last_event_ext = (
            (provider_connection.settings or {}).get("megafon_last_event_ext", "")
            if provider_connection
            else ""
        )
        telephony.megafon_last_history_user = (
            (provider_connection.settings or {}).get("megafon_last_history_user", "")
            if provider_connection
            else ""
        )
        telephony.megafon_last_history_ext = (
            (provider_connection.settings or {}).get("megafon_last_history_ext", "")
            if provider_connection
            else ""
        )
        telephony.megafon_last_history_status = (
            (provider_connection.settings or {}).get("megafon_last_history_status", "")
            if provider_connection
            else ""
        )
    return render(
        request,
        "pool_service/communications/channels.html",
        {
            "active_tab": "communications",
            "providers": providers,
            "telephony_connections": telephony_connections,
        },
    )


@login_required
@sensitive_post_parameters("client_id", "client_secret")
@transaction.atomic
def communication_connection_create(request, kind):
    organization = _context(request, "can_manage_channels")
    if kind not in {CommunicationChannel.KIND_WEBSITE, CommunicationChannel.KIND_AVITO}:
        raise Http404

    form = CommunicationConnectionForm(
        request.POST or None,
        kind=kind,
        require_avito_credentials=(kind == CommunicationChannel.KIND_AVITO),
    )
    if request.method == "POST" and form.is_valid():
        channel = _communication_provider_channel(organization, kind)
        external_id = form.cleaned_data["external_id"].strip()
        if ChannelConnection.objects.filter(channel=channel, external_id=external_id).exists():
            form.add_error("external_id", "Подключение с таким идентификатором уже существует.")
        else:
            connection = ChannelConnection.objects.create(
                channel=channel,
                name=form.cleaned_data["name"].strip(),
                external_id=external_id,
                is_active=form.cleaned_data["is_active"],
            )
            if kind == CommunicationChannel.KIND_AVITO:
                AvitoCredential.objects.create(
                    connection=connection,
                    client_id_encrypted=encrypt_secret(form.cleaned_data["client_id"]),
                    client_secret_encrypted=encrypt_secret(form.cleaned_data["client_secret"]),
                )
                connection.settings = {"avito_webhook_status": "not_connected"}
                connection.save(update_fields=["settings"])
                messages.success(
                    request,
                    "Аккаунт Авито сохранён. Нажмите «Подключить Авито», чтобы Service2 зарегистрировал webhook автоматически.",
                )
                return redirect("communications_channels")
            token = secrets.token_urlsafe(32)
            connection.set_api_token(token)
            connection.save(update_fields=["api_token_hash"])
            return _connection_setup_result(
                request, connection, token, secret_kind="website"
            )

    return render(
        request,
        "pool_service/communications/connection_form.html",
        {
            "active_tab": "communications",
            "form": form,
            "kind": kind,
            "provider_name": dict(CommunicationChannel.KIND_CHOICES)[kind],
            "editing": False,
        },
    )


@login_required
@sensitive_post_parameters("client_id", "client_secret")
@never_cache
@transaction.atomic
def communication_connection_edit(request, connection_id):
    from pool_service.avito_status_monitor import display_state
    organization = _context(request, "can_manage_channels")
    connection = get_object_or_404(
        ChannelConnection.objects.select_for_update().select_related("channel"),
        pk=connection_id,
        channel__organization=organization,
        channel__kind__in=(
            CommunicationChannel.KIND_WEBSITE,
            CommunicationChannel.KIND_AVITO,
        ),
    )
    kind = connection.channel.kind
    initial = {
        "name": connection.name,
        "external_id": connection.external_id,
        "is_active": connection.is_active,
    }
    form = CommunicationConnectionForm(
        request.POST or None,
        kind=kind,
        require_avito_credentials=False,
        initial=initial,
    )
    if request.method == "POST" and form.is_valid():
        external_id = form.cleaned_data["external_id"].strip()
        duplicate = ChannelConnection.objects.filter(
            channel=connection.channel,
            external_id=external_id,
        ).exclude(pk=connection.pk).exists()
        if duplicate:
            form.add_error("external_id", "Подключение с таким идентификатором уже существует.")
        else:
            previous_external_id = connection.external_id
            connection.name = form.cleaned_data["name"].strip()
            connection.external_id = external_id
            connection.is_active = form.cleaned_data["is_active"]
            connection.save(update_fields=["name", "external_id", "is_active"])
            avito_credentials_changed = False
            if kind == CommunicationChannel.KIND_AVITO and form.cleaned_data.get("client_id"):
                credential, _ = AvitoCredential.objects.get_or_create(
                    connection=connection,
                    defaults={
                        "client_id_encrypted": encrypt_secret(form.cleaned_data["client_id"]),
                        "client_secret_encrypted": encrypt_secret(form.cleaned_data["client_secret"]),
                    },
                )
                credential.client_id_encrypted = encrypt_secret(form.cleaned_data["client_id"])
                credential.client_secret_encrypted = encrypt_secret(form.cleaned_data["client_secret"])
                credential.access_token_encrypted = ""
                credential.access_token_expires_at = None
                credential.save()
                avito_credentials_changed = True
            if kind == CommunicationChannel.KIND_AVITO and (
                previous_external_id != external_id or avito_credentials_changed
            ):
                settings_data = dict(connection.settings or {})
                settings_data["avito_webhook_status"] = "needs_check"
                settings_data["avito_webhook_checked_at"] = ""
                connection.settings = settings_data
                connection.save(update_fields=["settings"])
            messages.success(request, "Настройки подключения сохранены.")
            return redirect("communications_channels")

    avito_credentials_configured = (
        kind == CommunicationChannel.KIND_AVITO
        and AvitoCredential.objects.filter(connection=connection).exists()
    )
    avito_diagnostic_accounts = []
    if kind == CommunicationChannel.KIND_AVITO:
        from pool_service.avito_workspace import diagnostic_account
        avito_diagnostic_accounts = [diagnostic_account(connection, has_credentials=avito_credentials_configured)]

    return render(
        request,
        "pool_service/communications/connection_form.html",
        {
            "active_tab": "communications",
            "form": form,
            "kind": kind,
            "provider_name": connection.channel.get_kind_display(),
            "editing": True,
            "connection": connection,
            "avito_credentials_configured": avito_credentials_configured,
            "avito_diagnostic_accounts": avito_diagnostic_accounts,
            "avito_timezone": getattr(settings, "COMMUNICATION_TIME_ZONE", "UTC"),
            "status_monitor_state": display_state(connection, request.user) if kind == CommunicationChannel.KIND_AVITO else {},
        },
    )


@login_required
@require_POST
@transaction.atomic
def communication_connection_rotate_token(request, connection_id):
    organization = _context(request, "can_manage_channels")
    connection = get_object_or_404(
        ChannelConnection.objects.select_for_update().select_related("channel"),
        pk=connection_id,
        channel__organization=organization,
        channel__kind__in=(
            CommunicationChannel.KIND_WEBSITE,
            CommunicationChannel.KIND_AVITO,
        ),
    )
    if connection.channel.kind == CommunicationChannel.KIND_AVITO:
        messages.info(
            request,
            "Webhook Авито теперь подключается и обновляется автоматически кнопкой «Подключить Авито».",
        )
        return redirect("communications_channels")
    token = secrets.token_urlsafe(32)
    connection.set_api_token(token)
    connection.save(update_fields=["api_token_hash"])
    return _connection_setup_result(
        request,
        connection,
        token,
        secret_kind=connection.channel.kind,
    )


def _avito_callback_url(request, connection, token):
    callback_url = request.build_absolute_uri(
        reverse("avito_webhook", args=[connection.public_id, token])
    )
    if not callback_url.startswith("https://"):
        raise AvitoError("service2_webhook_requires_https")
    return callback_url


def _avito_subscription_token(request, connection, value):
    if not isinstance(value, str):
        return None
    try:
        parsed = urlsplit(value)
    except ValueError:
        return None
    expected_host = request.get_host().lower()
    if (
        parsed.scheme != "https"
        or parsed.netloc.lower() != expected_host
        or parsed.query
        or parsed.fragment
    ):
        return None
    prefix = f"/api/communications/avito/{connection.public_id}/"
    suffix = "/webhook/"
    if not parsed.path.startswith(prefix) or not parsed.path.endswith(suffix):
        return None
    token = parsed.path[len(prefix):-len(suffix)]
    return token or None


def _set_avito_webhook_status(connection, status, *, error=""):
    with transaction.atomic():
        locked = ChannelConnection.objects.select_for_update().get(pk=connection.pk)
        settings_data = dict(locked.settings or {})
        settings_data["avito_webhook_status"] = status
        settings_data["avito_webhook_checked_at"] = timezone.now().isoformat()
        if error:
            settings_data["avito_webhook_error"] = error[:120]
        else:
            settings_data.pop("avito_webhook_error", None)
        locked.settings = settings_data
        locked.save(update_fields=["settings"])
        connection.settings = settings_data


@login_required
@require_POST
def communication_avito_connect(request, connection_id):
    organization = _context(request, "can_manage_channels")
    connection = get_object_or_404(
        ChannelConnection.objects.select_related("channel"),
        pk=connection_id,
        channel__organization=organization,
        channel__kind=CommunicationChannel.KIND_AVITO,
    )
    if not connection.is_active or not connection.channel.is_active:
        messages.error(request, "Сначала включите канал и подключение Авито.")
        return redirect("communications_channels")
    if not AvitoCredential.objects.filter(connection=connection).exists():
        messages.error(request, "Сначала сохраните client_id и client_secret Авито.")
        return redirect("communication_connection_edit", connection_id=connection.pk)

    force_reconnect = request.POST.get("force") == "1"
    try:
        provider_account_id = avito_authorized_account_id(connection)
        if provider_account_id != str(connection.external_id):
            duplicate = ChannelConnection.objects.filter(
                channel=connection.channel,
                external_id=provider_account_id,
            ).exclude(pk=connection.pk).exists()
            if duplicate:
                raise AvitoError("provider_account_already_connected")
            previous_account_id = connection.external_id
            connection.external_id = provider_account_id
            connection.save(update_fields=["external_id"])
            messages.info(
                request,
                f"Service2 исправил ID Авито: {previous_account_id} → {provider_account_id}.",
            )
        avito_verify_messenger_access(connection, provider_account_id)
    except AvitoError as exc:
        _set_avito_webhook_status(connection, "error", error=str(exc))
        if str(exc) in {"provider_http_402", "provider_http_403"}:
            messages.error(
                request,
                "Avito Messenger API недоступен для этих ключей. Проверьте выданный доступ/тариф Messenger API.",
            )
        else:
            messages.error(
                request,
                "Не удалось подтвердить аккаунт Авито и доступ Messenger API. Проверьте ключи и повторите подключение.",
            )
        return redirect("communications_channels")

    try:
        existing_urls = avito_webhook_subscriptions(connection)
    except AvitoError:
        # A temporary subscription-list failure must not block a fresh
        # registration; the new subscription is verified after registration.
        existing_urls = []

    if not force_reconnect:
        for value in existing_urls:
            token = _avito_subscription_token(request, connection, value)
            if token and connection.check_api_token(token):
                _set_avito_webhook_status(connection, "connected")
                messages.success(request, "Авито уже подключено. Webhook подтверждён.")
                return redirect("communications_channels")

    token = secrets.token_urlsafe(32)
    try:
        callback_url = _avito_callback_url(request, connection, token)
    except AvitoError as exc:
        _set_avito_webhook_status(connection, "error", error=str(exc))
        messages.error(request, "Service2 не смог сформировать защищённый HTTPS webhook URL.")
        return redirect("communications_channels")

    # Commit the new webhook hash before calling Avito. Avito can probe the
    # callback synchronously during subscription, so a transaction spanning the
    # provider call would make the new secret invisible to that probe.
    with transaction.atomic():
        locked = ChannelConnection.objects.select_for_update().get(pk=connection.pk)
        old_hash = locked.api_token_hash
        locked.set_api_token(token)
        locked.save(update_fields=["api_token_hash"])
        connection.api_token_hash = locked.api_token_hash

    registration_error = None
    try:
        avito_subscribe_webhook(connection, callback_url)
    except AvitoError as exc:
        # The provider may have accepted the subscription even when its response
        # was lost. Verify the authoritative subscription list before rollback.
        registration_error = exc

    try:
        verified_urls = avito_webhook_subscriptions(connection)
        verified = any(
            _avito_subscription_token(request, connection, value) == token
            for value in verified_urls
        )
    except AvitoError as exc:
        verified_urls = []
        verified = False
        if registration_error is None:
            registration_error = exc

    if not verified:
        rolled_back = False
        with transaction.atomic():
            locked = ChannelConnection.objects.select_for_update().get(pk=connection.pk)
            # Do not clobber a newer concurrent reconnect.
            if locked.check_api_token(token):
                locked.api_token_hash = old_hash
                locked.save(update_fields=["api_token_hash"])
                connection.api_token_hash = old_hash
                rolled_back = True
        if rolled_back:
            error = registration_error or AvitoError("provider_webhook_not_confirmed")
            _set_avito_webhook_status(connection, "error", error=str(error))
            messages.error(
                request,
                "Авито не подтвердило webhook. Проверьте client_id, client_secret, доступ Messenger API и повторите подключение.",
            )
        else:
            messages.warning(
                request,
                "Подключение Авито уже было изменено другим запросом. Обновите страницу и проверьте статус.",
            )
        return redirect("communications_channels")

    for stale_url in existing_urls:
        if stale_url == callback_url:
            continue
        if _avito_subscription_token(request, connection, stale_url):
            try:
                avito_unsubscribe_webhook(connection, stale_url)
            except AvitoError:
                pass

    _set_avito_webhook_status(connection, "connected")
    messages.success(request, "Авито подключено: webhook зарегистрирован и подтверждён.")
    return redirect("communications_channels")


@login_required
@require_POST
def communication_avito_check(request, connection_id):
    organization = _context(request, "can_manage_channels")
    connection = get_object_or_404(
        ChannelConnection.objects.select_related("channel"),
        pk=connection_id,
        channel__organization=organization,
        channel__kind=CommunicationChannel.KIND_AVITO,
    )
    try:
        provider_account_id = avito_authorized_account_id(connection)
        if provider_account_id != str(connection.external_id):
            duplicate = ChannelConnection.objects.filter(
                channel=connection.channel,
                external_id=provider_account_id,
            ).exclude(pk=connection.pk).exists()
            if duplicate:
                raise AvitoError("provider_account_already_connected")
            connection.external_id = provider_account_id
            connection.save(update_fields=["external_id"])
        avito_verify_messenger_access(connection, provider_account_id)
        urls = avito_webhook_subscriptions(connection)
        connected = any(
            token and connection.check_api_token(token)
            for token in (
                _avito_subscription_token(request, connection, value)
                for value in urls
            )
        )
    except AvitoError as exc:
        _set_avito_webhook_status(connection, "error", error=str(exc))
        if str(exc) in {"provider_http_402", "provider_http_403"}:
            messages.error(
                request,
                "Avito Messenger API недоступен для этих ключей. Проверьте выданный доступ/тариф Messenger API.",
            )
        else:
            messages.error(request, "Не удалось проверить подключение Авито.")
        return redirect("communications_channels")

    if connected:
        _set_avito_webhook_status(connection, "connected")
        messages.success(request, "Подключение Авито подтверждено.")
    else:
        _set_avito_webhook_status(connection, "not_connected")
        messages.warning(request, "Webhook Service2 не найден в активных подписках Авито.")
    return redirect("communications_channels")


@login_required
@require_POST
def communication_avito_sync(request, connection_id):
    organization = _context(request, "can_manage_channels")
    connection = get_object_or_404(
        ChannelConnection.objects.select_related("channel"),
        pk=connection_id,
        channel__organization=organization,
        channel__kind=CommunicationChannel.KIND_AVITO,
    )
    if not connection.is_active or not connection.channel.is_active:
        messages.error(request, "Сначала включите канал и подключение Авито.")
        return redirect("communications_channels")
    if not AvitoCredential.objects.filter(connection=connection).exists():
        messages.error(request, "Сначала сохраните client_id и client_secret Авито.")
        return redirect("communication_connection_edit", connection_id=connection.pk)

    try:
        result = avito_sync_recent_messages(connection)
    except AvitoError as exc:
        with transaction.atomic():
            locked = ChannelConnection.objects.select_for_update().get(pk=connection.pk)
            settings_data = dict(locked.settings or {})
            settings_data["avito_pull_last_checked_at"] = timezone.now().isoformat()
            settings_data["avito_pull_last_error"] = str(exc)[:120]
            locked.settings = settings_data
            locked.save(update_fields=["settings"])
        if str(exc) in {"provider_http_402", "provider_http_403"}:
            messages.error(
                request,
                "Avito Messenger API не даёт читать сообщения этими ключами. Нужен доступ/тариф Messenger API.",
            )
        else:
            messages.error(request, f"Не удалось синхронизировать сообщения Авито: {exc}")
        return redirect("communications_channels")

    with transaction.atomic():
        locked = ChannelConnection.objects.select_for_update().get(pk=connection.pk)
        settings_data = dict(locked.settings or {})
        settings_data["avito_pull_last_checked_at"] = timezone.now().isoformat()
        settings_data["avito_pull_last_created"] = result.messages_created
        settings_data["avito_pull_last_existing"] = result.messages_existing
        settings_data["avito_pull_last_checked_messages"] = result.messages_checked
        settings_data["avito_pull_last_chats"] = result.chats_checked
        settings_data.pop("avito_pull_last_error", None)
        locked.settings = settings_data
        locked.save(update_fields=["settings"])

    messages.success(
        request,
        (
            f"Синхронизация Авито завершена: новых сообщений {result.messages_created}, "
            f"уже были {result.messages_existing}, проверено чатов {result.chats_checked}."
        ),
    )
    return redirect("communications_conversations")


@login_required
@require_POST
@transaction.atomic
def communication_telephony_connect(request, connection_id):
    organization = _context(request, "can_manage_channels")
    telephony = get_object_or_404(
        TelephonyConnection.objects.select_for_update(),
        pk=connection_id,
        organization=organization,
    )
    provider_connection = _telephony_provider_connection(telephony)
    provider_settings = dict(provider_connection.settings or {})
    if not (
        provider_settings.get("megafon_api_base_url")
        and provider_settings.get("megafon_api_key_encrypted")
    ):
        messages.error(
            request,
            "Сначала скопируйте из МегаФона «Адрес АТС» и «Ключ для авторизации в АТС».",
        )
        return redirect("communication_telephony_edit", connection_id=telephony.pk)
    token = secrets.token_urlsafe(32)
    provider_connection.set_api_token(token)
    provider_connection.is_active = True
    provider_connection.save(update_fields=["api_token_hash", "is_active"])
    return _telephony_setup_result(
        request,
        telephony,
        provider_connection,
        token,
    )


@login_required
@sensitive_post_parameters("ats_api_key")
@transaction.atomic
def communication_telephony_create(request):
    organization = _context(request, "can_manage_channels")
    form = TelephonyConnectionForm(
        request.POST or None,
        require_ats_api_key=True,
    )
    if request.method == "POST" and form.is_valid():
        external_id = form.cleaned_data["external_id"].strip()
        if TelephonyConnection.objects.filter(
            organization=organization, external_id=external_id
        ).exists():
            form.add_error("external_id", "Линия с таким идентификатором уже существует.")
        else:
            telephony = TelephonyConnection.objects.create(
                organization=organization,
                name=form.cleaned_data["name"].strip(),
                external_id=external_id,
                recording_allowed_hosts=form.cleaned_data["recording_allowed_hosts"],
                is_active=form.cleaned_data["is_active"],
            )
            provider_connection = _telephony_provider_connection(telephony)
            _save_megafon_api_settings(provider_connection, form)
            messages.success(
                request,
                "Данные АТС МегаФона сохранены. Теперь настройте обратную отправку событий в Service2.",
            )
            return redirect("communications_channels")
    return render(
        request,
        "pool_service/communications/telephony_form.html",
        {"active_tab": "communications", "form": form, "editing": False},
    )


@login_required
@sensitive_post_parameters("ats_api_key")
@transaction.atomic
def communication_telephony_edit(request, connection_id):
    organization = _context(request, "can_manage_channels")
    connection = get_object_or_404(
        TelephonyConnection.objects.select_for_update(),
        pk=connection_id,
        organization=organization,
    )
    provider_connection = _telephony_provider_connection(connection)
    provider_settings = dict(provider_connection.settings or {})
    initial = {
        "name": connection.name,
        "external_id": connection.external_id,
        "ats_base_url": provider_settings.get("megafon_api_base_url", ""),
        "recording_allowed_hosts": "\n".join(connection.recording_allowed_hosts or []),
        "is_active": connection.is_active,
    }
    form = TelephonyConnectionForm(
        request.POST or None,
        initial=initial,
        require_ats_api_key=not bool(
            provider_settings.get("megafon_api_key_encrypted")
        ),
    )
    if request.method == "POST" and form.is_valid():
        external_id = form.cleaned_data["external_id"].strip()
        duplicate = TelephonyConnection.objects.filter(
            organization=organization, external_id=external_id
        ).exclude(pk=connection.pk).exists()
        if duplicate:
            form.add_error("external_id", "Линия с таким идентификатором уже существует.")
        else:
            previous_external_id = connection.external_id
            connection.name = form.cleaned_data["name"].strip()
            connection.external_id = external_id
            connection.recording_allowed_hosts = form.cleaned_data["recording_allowed_hosts"]
            connection.is_active = form.cleaned_data["is_active"]
            connection.save(
                update_fields=[
                    "name",
                    "external_id",
                    "recording_allowed_hosts",
                    "is_active",
                ]
            )
            if previous_external_id != external_id:
                provider_connection = ChannelConnection.objects.filter(
                    channel__organization=organization,
                    channel__kind=CommunicationChannel.KIND_MEGAFON,
                    external_id=previous_external_id,
                ).first()
                if provider_connection and not ChannelConnection.objects.filter(
                    channel=provider_connection.channel,
                    external_id=external_id,
                ).exclude(pk=provider_connection.pk).exists():
                    provider_connection.external_id = external_id
                    provider_connection.name = connection.name
                    provider_connection.is_active = connection.is_active
                    provider_connection.save(
                        update_fields=["external_id", "name", "is_active"]
                    )
            provider_connection = _telephony_provider_connection(connection)
            _save_megafon_api_settings(provider_connection, form)
            messages.success(request, "Настройки МегаФона сохранены.")
            return redirect("communications_channels")
    return render(
        request,
        "pool_service/communications/telephony_form.html",
        {
            "active_tab": "communications",
            "form": form,
            "editing": True,
            "connection": connection,
        },
    )


@login_required
@require_POST
@transaction.atomic
def communication_telephony_set_active(request, connection_id):
    organization = _context(request, "can_manage_channels")
    desired = _requested_active(request)
    if desired is None:
        return HttpResponseBadRequest("Некорректный статус подключения.")
    connection = get_object_or_404(
        TelephonyConnection.objects.select_for_update(),
        pk=connection_id,
        organization=organization,
    )
    if connection.is_active != desired:
        connection.is_active = desired
        connection.save(update_fields=["is_active"])
    provider_connection = _telephony_provider_connection(connection)
    if provider_connection.is_active != desired:
        provider_connection.is_active = desired
        provider_connection.save(update_fields=["is_active"])
    messages.success(
        request,
        "Подключение Мегафона включено." if desired else "Подключение Мегафона отключено.",
    )
    return redirect("communications_channels")


def _requested_active(request):
    value = request.POST.get("active")
    if value not in {"0", "1"}:
        return None
    return value == "1"


@login_required
@require_POST
@transaction.atomic
def communication_channel_set_active(request, channel_id):
    organization = _context(request, "can_manage_channels")
    desired = _requested_active(request)
    if desired is None:
        return HttpResponseBadRequest("Некорректный статус канала.")
    channel = get_object_or_404(
        CommunicationChannel.objects.select_for_update(),
        pk=channel_id,
        organization=organization,
    )
    if channel.is_active != desired:
        channel.is_active = desired
        channel.save(update_fields=["is_active"])
    messages.success(request, "Канал включён." if desired else "Канал отключён.")
    return redirect("communications_channels")


@login_required
@require_POST
@transaction.atomic
def communication_connection_set_active(request, connection_id):
    organization = _context(request, "can_manage_channels")
    desired = _requested_active(request)
    if desired is None:
        return HttpResponseBadRequest("Некорректный статус подключения.")
    connection = get_object_or_404(
        ChannelConnection.objects.select_for_update().select_related("channel"),
        pk=connection_id,
        channel__organization=organization,
    )
    if connection.is_active != desired:
        connection.is_active = desired
        connection.save(update_fields=["is_active"])
    messages.success(request, "Подключение включено." if desired else "Подключение отключено.")
    return redirect("communications_channels")


@login_required
@ensure_csrf_cookie
def communication_notification_feed(request):
    organization = _context(request, "can_view_conversations")
    visible_ids = []
    raw_visible = request.GET.get("visible", "")
    if raw_visible:
        parts = raw_visible.split(",")
        if (
            len(raw_visible) > 500
            or len(parts) > 50
            or any(not item.isascii() or not item.isdecimal() for item in parts)
        ):
            return JsonResponse({"error": "invalid_visible"}, status=400)
        visible_ids = [int(item) for item in parts]
        if any(item <= 0 or item > 9_223_372_036_854_775_807 for item in visible_ids):
            return JsonResponse({"error": "invalid_visible"}, status=400)
    notifications = list(Notification.objects.filter(
        user=request.user,
        organization=organization,
        kind="communication",
        is_resolved=False,
    ).order_by("-created_at", "-pk")[:10])
    result = []
    for notification in reversed(notifications):
        action_url = notification.action_url
        if not action_url.startswith("/communications/") or action_url.startswith("//"):
            action_url = reverse("communications_conversations")
        result.append({
            "id": notification.pk,
            "title": notification.title,
            "message": notification.message,
            "action_url": action_url,
            "created_at": notification.created_at.isoformat(),
        })
    resolved_ids = list(Notification.objects.filter(
        user=request.user,
        organization=organization,
        kind="communication",
        is_resolved=True,
        pk__in=visible_ids,
    ).values_list("pk", flat=True))
    return JsonResponse({"notifications": result, "resolved_ids": resolved_ids})


@login_required
@require_POST
def communication_notification_resolve(request, notification_id):
    organization = _context(request, "can_view_conversations")
    notification = get_object_or_404(
        Notification,
        pk=notification_id,
        user=request.user,
        organization=organization,
        kind="communication",
        is_resolved=False,
    )
    notification.is_read = True
    notification.is_resolved = True
    notification.resolved_at = timezone.now()
    notification.save(update_fields=["is_read", "is_resolved", "resolved_at"])
    return JsonResponse({"resolved": True})
