"""Preserve Avito's XLSX category fields; only edit requested content fields.

Imports and drafts never change the publicly served published snapshot.
No provider URL is fetched from a workbook and no raw workbook is persisted.
"""
from copy import deepcopy
from io import BytesIO
import re
import time
import hashlib
from urllib.parse import urlsplit
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from openpyxl.utils.exceptions import InvalidFileException

from openpyxl import load_workbook, Workbook
from openpyxl.utils.cell import column_index_from_string
from django import forms
from django.core.paginator import Paginator
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from django.urls import reverse
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_GET, require_http_methods

from .avito_feed_models import AvitoServicesFeed
from .avito_services_history import ensure_baseline, record_version, require_identities, restored_snapshot
from .avito_services_xml import xml_bytes
from .avito_publication_models import AvitoServicesPublication, AvitoServicesListingIntent
from .avito_services_publication import bind_publication, publication_context, archive_context, _binding_valid, public_origin, annotate_reactivation
from .avito_services_media import (validate_photo_scope, validate_snapshot_media,
                                  mark_snapshot_media_published, editor_media_context)
from .avito_management import _scope
from .communication_models import ChannelConnection

MAX_BYTES = 4 * 1024 * 1024
MAX_ROWS = 200
MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def images(value, *, max_count=10):
    urls = [part.strip() for part in re.split(r'[|\r\n]+', value) if part.strip()]
    if max_count not in (10, 30) or not 1 <= len(urls) <= max_count:
        raise ValueError(f'Укажите от 1 до {max_count} ссылок на фотографии.')
    for url in urls:
        try:
            parsed = urlsplit(url)
            valid = parsed.scheme in {'https', 'http'} and parsed.hostname and not parsed.username and not parsed.password
        except ValueError:
            valid = False
        if not valid or len(url) > 2000 or any(c.isspace() for c in url):
            raise ValueError('Некорректная ссылка на фотографию.')
    result = ' | '.join(urls)
    if len(result) > 32767:
        raise ValueError('Список фотографий длиннее 32767 символов.')
    return result


def parse_export(raw):
    if len(raw) > MAX_BYTES:
        raise ValueError('Файл больше 4 МБ.')
    try:
        with ZipFile(BytesIO(raw)) as archive:
            infos = archive.infolist()
            if len(infos) > 500 or sum(i.file_size for i in infos) > 20 * MAX_BYTES:
                raise ValueError('Превышен допустимый размер распакованного файла.')
            # Reject DTD/entity expansion even when optional XML hardening is absent.
            for info in infos:
                if info.filename.endswith('.xml'):
                    content = archive.read(info).replace(b'\x00', b'').upper()
                    if info.filename.startswith('xl/worksheets/'):
                        for column, row in re.findall(rb'<(?:[A-Z0-9_]+:)?C\b[^>]*\bR=[\"\']([A-Z]+)([0-9]+)[\"\']', content):
                            if len(column) > 3 or column_index_from_string(column.decode()) > 150 or int(row) > 2000:
                                raise ValueError('Превышены размеры листа.')
                    if b'<!DOCTYPE' in content or b'<!ENTITY' in content:
                        raise ValueError('XML с DTD не поддерживается.')
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        result, ids, avito_ids = [], set(), set()
        if len(book.sheetnames) > 50:
            book.close()
            raise ValueError('Слишком много листов.')
        total_cells = 0
        try:
            for sheet in book:
                # Avito reverse exports may declare A1:A1 despite containing full
                # sheets. Ignore the dimension hint, but bound actual iteration.
                sheet.reset_dimensions()
                rows = list(sheet.iter_rows(max_row=2001, max_col=151, values_only=True))
                while rows and not any(v is not None for v in rows[-1]):
                    rows.pop()
                total_cells += sum(sum(v is not None for v in row) for row in rows)
                if total_cells > 100000 or len(rows) > 2000 or any(row[150] is not None for row in rows):
                    raise ValueError('Превышены размеры листа.')
                width = max((i + 1 for row in rows for i, v in enumerate(row) if v is not None), default=0)
                rows = [row[:width] for row in rows]
                if len(rows) < 5 or 'AvitoId' not in rows[1]:
                    continue
                headers = [str(h) if h is not None else '' for h in rows[1]]
                named = [h for h in headers if h]
                if len(named) != len(set(named)) or any(not re.fullmatch(r'[A-Za-z][A-Za-z0-9]*', h) for h in named):
                    raise ValueError('Некорректные заголовки столбцов.')
                entries = []
                for row in rows[4:]:
                    if not any(v is not None for v in row):
                        continue
                    data = {h: str(v) if v is not None else '' for h, v in zip(headers, row) if h}
                    if any(len(v) > 32767 for v in data.values()):
                        raise ValueError('Значение ячейки длиннее 32767 символов.')
                    if any(v.startswith('=') for v in data.values()):
                        raise ValueError('Формулы в данных объявлений не поддерживаются.')
                    if data.get('Category') != 'Предложение услуг' or data.get('AvitoStatus') != 'Активно':
                        raise ValueError('Импорт поддерживает только активные объявления услуг.')
                    if not re.fullmatch(r'[0-9]{1,20}', data.get('AvitoId', '')) or not data.get('Id'):
                        raise ValueError('В каждой строке нужны Id и AvitoId существующего объявления.')
                    if data['Id'] in ids or data['AvitoId'] in avito_ids:
                        raise ValueError('Повторяется Id или AvitoId.')
                    ids.add(data['Id']); avito_ids.add(data['AvitoId'])
                    if len(ids) > MAX_ROWS:
                        raise ValueError('Поддерживается не более 200 объявлений.')
                    if not data.get('Title') or not data.get('Description'):
                        raise ValueError('Отсутствует заголовок или описание.')
                    data['ImageUrls'] = images(data.get('ImageUrls', ''))
                    if data.get('ImageNames'):
                        raise ValueError('Используйте ссылки на фото, а не локальные имена файлов.')
                    if data.get('ListingFee') not in {'', 'Package'}:
                        raise ValueError('Разрешено только размещение из пакета (Package).')
                    entries.append(data)
                if entries:
                    result.append({'name': sheet.title, 'intro': str(rows[0][0] or ''),
                                   'headers': headers, 'required': [str(v) if v is not None else None for v in rows[2]],
                                   'hints': [str(v) if v is not None else None for v in rows[3]], 'rows': entries})
        finally:
            book.close()
    except (BadZipFile, KeyError, TypeError, OSError, ParseError, InvalidFileException, IndexError) as error:
        raise ValueError('Не удалось прочитать экспорт Авито XLSX.') from error
    if not result:
        raise ValueError('В файле нет активных объявлений услуг.')
    return result


def workbook_bytes(sheets):
    if any(len(str(value)) > 32767 for row in entries(sheets) for value in row.values()):
        raise ValueError('Значение ячейки длиннее 32767 символов.')
    book = Workbook()
    book.remove(book.active)
    for data in sheets:
        sheet = book.create_sheet(data['name'])
        sheet.append([data['intro']])
        sheet.append(data['headers'])
        sheet.append(data['required'])
        sheet.append(data['hints'])
        for entry in data['rows']:
            sheet.append([entry.get(h, '') for h in data['headers']])
        # User strings must never become formulas or external workbook links.
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    output = BytesIO()
    book.save(output)
    return output.getvalue()


class ContentForm(forms.Form):
    title = forms.CharField(label='Заголовок', max_length=100)
    description = forms.CharField(label='Описание', max_length=32767, widget=forms.Textarea(attrs={'rows': 12}))
    image_urls = forms.CharField(label='Фотографии: одна ссылка на строку, максимум 10', max_length=32767, widget=forms.Textarea(attrs={'rows': 7}))
    reason = forms.CharField(label='Причина изменения', max_length=500, required=False)

    def clean_image_urls(self):
        try:
            return images(self.cleaned_data['image_urls'])
        except ValueError as error:
            raise forms.ValidationError(str(error)) from error


def connection_for(request, connection_id):
    return get_object_or_404(ChannelConnection, pk=connection_id, channel__organization=_scope(request),
                            channel__kind='avito', is_active=True, channel__is_active=True)


def bound(feed, connection):
    return feed.account_id == connection.external_id and feed.organization_id_at_import == connection.channel.organization_id


def entries(sheets):
    return [row for sheet in sheets for row in sheet['rows']]


def _reason(request, default):
    value = request.POST.get('reason', '').strip()
    if len(value) > 500:
        raise ValueError('Причина изменения длиннее 500 символов.')
    return value or default


def _version(feed, value):
    if (not isinstance(value, str) or not re.fullmatch(r'[0-9]{1,19}', value)
            or not 0 < int(value) <= 2**63 - 1):
        raise Http404
    return get_object_or_404(feed.versions, pk=value)


def _page_context(request, connection, feed, error=''):
    observations = {row.get('avito_id'): row for row in feed.latest_observations} if feed else {}
    publication = publication_context(feed) if feed else None
    reported = {row['avito_id']: row for row in publication['rows']} if publication else {}
    intents = {intent.avito_id: intent for intent in feed.listing_intents.filter(
        account_id=connection.external_id, organization_id=connection.channel.organization_id)} if feed else {}
    rows = [{**row, 'observation': observations.get(row['AvitoId']),
             'archive_control': archive_context(row, observations.get(row['AvitoId']),
                 intents.get(row['AvitoId']), reported.get(row['AvitoId']))}
            for row in entries(feed.draft)] if feed else []
    history = Paginator(feed.versions.select_related('actor').order_by('-pk'), 20).get_page(
        request.GET.get('history_page')) if feed else None
    selected = None
    if feed and request.GET.get('history_version'):
        selected = _version(feed, request.GET['history_version'])
    return {
        'connection': connection, 'feed': feed, 'rows': rows, 'error': error,
        'history': history, 'selected_version': selected, 'publication': publication,
        'feed_url': request.build_absolute_uri(reverse('avito_services_feed_public', args=[feed.public_key]))
                    if feed and feed.enabled else '',
    }


@login_required
@never_cache
@require_http_methods(['GET', 'POST'])
def services_feed(request, connection_id):
    connection = connection_for(request, connection_id)
    feed = AvitoServicesFeed.objects.filter(connection=connection).first()
    error = ''
    if feed and not bound(feed, connection):
        return HttpResponse('Аккаунт подключения изменился. Старый источник заблокирован.', status=409)
    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'verify' and feed:
            return _verify_publication(request, connection, feed)
        try:
            imported = None
            if action == 'import':
                upload = request.FILES.get('file')
                if not upload or upload.size > MAX_BYTES:
                    raise ValueError('Выберите экспорт Авито XLSX размером до 4 МБ.')
                if not upload.name.startswith(connection.external_id + '_'):
                    raise ValueError('Имя экспорта должно начинаться с ID выбранного аккаунта и подчёркивания.')
                imported = parse_export(upload.read(MAX_BYTES + 1))
                for item in entries(imported):
                    item['ImageUrls'] = validate_photo_scope(connection, item['ImageUrls'])
            reason = _reason(request, {'import': 'Импорт обратной выгрузки Авито',
                                      'restore': 'Восстановление содержания прежней версии',
                                      'disable': 'Отключение доступа к файлу',
                                      'policy': 'Политика возврата объявления в Service2'}.get(action, ''))
            with transaction.atomic():
                connection = ChannelConnection.objects.select_for_update().get(pk=connection.pk)
                connection = connection_for(request, connection_id)
                feed = AvitoServicesFeed.objects.select_for_update().filter(connection=connection).first()
                if feed and (not bound(feed, connection) or str(feed.revision) != request.POST.get('revision')):
                    raise ValueError('Данные изменились. Обновите страницу перед сохранением.')
                source = None
                policy_changes = []
                before = deepcopy(feed.draft) if feed else []
                if action == 'import':
                    if feed:
                        require_identities(feed, imported)
                        ensure_baseline(feed)
                        feed.draft = imported
                    else:
                        feed = AvitoServicesFeed(connection=connection, account_id=connection.external_id,
                            organization_id_at_import=connection.channel.organization_id, draft=imported)
                elif action == 'restore' and feed:
                    source = _version(feed, request.POST.get('version_id'))
                    candidate = restored_snapshot(feed, source)
                    # A restore is a new draft, never a replacement of the public file.
                    xml_bytes(candidate, max_images=30)
                    for item in entries(candidate):
                        validate_photo_scope(connection, item['ImageUrls'], max_count=30)
                    require_identities(feed, candidate)
                    ensure_baseline(feed)
                    feed.draft = candidate
                elif action == 'policy' and feed:
                    avito_id, stopped = request.POST.get('avito_id'), request.POST.get('owner_withdrawn')
                    target = next((item for item in entries(feed.draft) if item['AvitoId'] == avito_id), None)
                    if target is None or stopped not in {'0', '1'}:
                        raise ValueError('Некорректная политика возврата объявления.')
                    ensure_baseline(feed)
                    intent, _ = AvitoServicesListingIntent.objects.get_or_create(feed=feed, avito_id=avito_id,
                        defaults={'account_id': feed.account_id, 'organization_id': feed.organization_id_at_import})
                    if intent.account_id != feed.account_id or intent.organization_id != feed.organization_id_at_import:
                        raise ValueError('Политика относится к другой привязке аккаунта. Нужна проверка источника.')
                    was_stopped = intent.owner_withdrawn
                    intent.owner_withdrawn, intent.marked_by = stopped == '1', request.user
                    intent.save()
                    policy_changes = [{'avito_id': avito_id, 'id': target['Id'],
                        'field': 'Запрет автоматического возврата (Service2)', 'before': was_stopped, 'after': intent.owner_withdrawn}]
                elif action == 'disable' and feed:
                    ensure_baseline(feed)
                    feed.enabled = False
                else:
                    raise ValueError('Неизвестное действие.')
                feed.revision += 1
                feed.updated_by = request.user
                feed.save()
                record_version(feed, action, before, feed.draft, actor=request.user,
                               reason=reason, source_version=source, extra_changes=policy_changes)
            return redirect('avito_services_feed', connection_id=connection_id)
        except ValueError as exc:
            error = str(exc)
    return render(request, 'pool_service/avito/services_feed.html', _page_context(request, connection, feed, error))


@login_required
@never_cache
@require_http_methods(['GET', 'POST'])
def services_feed_edit(request, connection_id, avito_id):
    connection = connection_for(request, connection_id)
    feed = get_object_or_404(AvitoServicesFeed, connection=connection)
    if not bound(feed, connection):
        raise Http404
    if request.method == 'POST' and request.POST.get('revision') != str(feed.revision):
        return HttpResponse('Данные изменились. Откройте объявление заново перед редактированием.', status=409)
    row = next((r for r in entries(feed.draft) if r['AvitoId'] == avito_id), None)
    if row is None:
        raise Http404
    form = ContentForm(request.POST or None, initial={'title': row['Title'], 'description': row['Description'],
                                                    'image_urls': row['ImageUrls'].replace(' | ', '\n')})
    if request.method == 'POST' and form.is_valid():
        try:
            with transaction.atomic():
                ChannelConnection.objects.select_for_update().get(pk=connection.pk)
                connection = connection_for(request, connection_id)
                locked = AvitoServicesFeed.objects.select_for_update().get(pk=feed.pk)
                if not bound(locked, connection) or request.POST.get('revision') != str(locked.revision):
                    return HttpResponse('Данные изменились. Откройте объявление заново перед редактированием.', status=409)
                form.cleaned_data['image_urls'] = validate_photo_scope(connection, form.cleaned_data['image_urls'])
                before = deepcopy(locked.draft)
                updated = deepcopy(locked.draft)
                target = next(r for r in entries(updated) if r['AvitoId'] == avito_id)
                target.update(Title=form.cleaned_data['title'], Description=form.cleaned_data['description'],
                              ImageUrls=form.cleaned_data['image_urls'])
                require_identities(locked, updated)
                ensure_baseline(locked)
                locked.draft = updated
                locked.revision += 1
                locked.updated_by = request.user
                locked.save()
                record_version(locked, 'edit', before, updated, actor=request.user,
                               reason=form.cleaned_data['reason'] or 'Изменение текста и фотографий')
                return redirect('avito_services_feed', connection_id=connection_id)
        except ValueError as error:
            # These are local history/XML validators, not provider response text.
            form.add_error(None, str(error))
            return render(request, 'pool_service/avito/services_edit.html',
                          {'form': form, 'feed': feed, 'connection': connection, 'row': row,
                           **editor_media_context(connection, row)}, status=400)
    return render(request, 'pool_service/avito/services_edit.html', {'form': form, 'feed': feed, 'connection': connection, 'row': row,
                   **editor_media_context(connection, row)})


@login_required
@never_cache
@require_GET
def services_feed_download(request, connection_id):
    connection = connection_for(request, connection_id)
    feed = get_object_or_404(AvitoServicesFeed, connection=connection)
    if not bound(feed, connection):
        raise Http404
    response = HttpResponse(workbook_bytes(feed.draft), content_type=MIME)
    response['Content-Disposition'] = 'attachment; filename="avito-services-draft.xlsx"'
    return response


def _deadline(value):
    if not isinstance(value, str):
        return None
    try:
        result = parse_datetime(value)
    except (ValueError, TypeError):
        return None
    return result.isoformat() if result and timezone.is_aware(result) else None


def verify_inventory(connection, sheets):
    """Verify all active coverage and owned identities; never restore inactive ads."""
    from .communication_avito import access_token, authorized_account_id, AvitoError
    from .avito_workspace import _get, fetch_item_detail
    account_id = authorized_account_id(connection)
    if str(account_id) != connection.external_id:
        raise AvitoError('provider_account_mismatch')
    token = access_token(connection)
    deadline = time.monotonic() + 45
    active = set()
    for page in range(1, 4):
        if time.monotonic() >= deadline:
            raise AvitoError('provider_check_limit')
        response = _get(token, f'/core/v1/items?per_page=100&page={page}&status=active')
        resources = response.get('resources')
        if not isinstance(resources, list) or len(resources) > 100:
            raise AvitoError('provider_data_invalid')
        for item in resources:
            item_id = item.get('id') if isinstance(item, dict) else None
            if type(item_id) is not int or not 0 < item_id < 10**32 or str(item_id) in active:
                raise AvitoError('provider_data_invalid')
            active.add(str(item_id))
        if len(resources) < 100:
            break
    else:
        raise ValueError('Слишком много активных объявлений для проверки этого файла.')
    rows = entries(sheets)
    file_ids = {row['AvitoId'] for row in rows}
    if len(file_ids) != len(rows) or not active.issubset(file_ids):
        raise ValueError('В файле отсутствуют активные объявления аккаунта или повторяются AvitoId. Нужна полная выгрузка.')
    labels = {'active': 'Активно', 'old': 'Архив; причина не подтверждена',
              'removed': 'Снято; причина не подтверждена', 'rejected': 'Отклонено', 'blocked': 'Заблокировано'}
    observed = []
    checked_at = timezone.now().isoformat()
    for row in rows:
        if time.monotonic() >= deadline:
            raise AvitoError('provider_check_limit')
        detail = fetch_item_detail(token, account_id, row['AvitoId'])
        status = detail.get('status')
        if status not in labels:
            raise AvitoError('provider_data_invalid')
        if detail.get('autoload_item_id') not in {'', row['Id']}:
            raise ValueError('Привязка Id объявления изменилась. Обновите экспорт.')
        if (status == 'active') != (row['AvitoId'] in active):
            raise ValueError('Статусы изменились во время проверки. Повторите проверку.')
        finish = _deadline(detail.get('finish_time'))
        observed.append({'avito_id': row['AvitoId'], 'id': row['Id'], 'status': status,
            'status_label': labels[status], 'included': status == 'active',
            'exclusion_reason': '' if status == 'active' else 'Неактивное объявление не отправляется повторно без отдельной проверки активации.',
            'finish_time': finish, 'finish_time_source': 'avito_item_detail.finish_time' if finish else '',
            'source': 'avito_item_detail', 'checked_at': checked_at})
    published = deepcopy(sheets)
    for sheet in published:
        sheet['rows'] = [row for row in sheet['rows'] if row['AvitoId'] in active]
    published = [sheet for sheet in published if sheet['rows']]
    return {'sheets': published, 'observations': observed}


def _older_observation(feed, observations):
    if not feed.last_checked_at or not observations:
        return False
    started = parse_datetime(observations[0]['checked_at'])
    return started is None or timezone.is_naive(started) or feed.last_checked_at > started


def _safe_failure(error):
    from .communication_avito import AvitoError
    if isinstance(error, ValueError):
        return str(error)
    code = str(error)
    safe_code = code if re.fullmatch(r'provider_[a-z0-9_]{1,80}', code) else 'provider_check_failed'
    return f'Не удалось проверить объявления через API Авито ({safe_code}). Источник не изменён.'


class PublicationConflict(ValueError):
    pass


def _publication_connection(connection_id, actor, *, locked=False):
    """Current actor/org permissions, both before I/O and under publication locks."""
    from django.contrib.auth import get_user_model
    from django.core.exceptions import PermissionDenied
    from .communication_services import organization_access, conversation_capability
    current_actor = get_user_model().objects.filter(pk=getattr(actor, 'pk', None), is_active=True).first()
    query = ChannelConnection.objects.select_related('channel', 'avito_credential')
    if locked:
        query = query.select_for_update()
    connection = query.filter(pk=connection_id, is_active=True, channel__is_active=True,
                              channel__kind='avito').first()
    if (current_actor is None or connection is None
            or organization_access(current_actor, connection.channel.organization) is None
            or not conversation_capability(current_actor, 'can_manage_channels', connection.channel.organization)):
        raise PermissionDenied
    return connection, current_actor


def publish_services_snapshot(connection, actor, *, expected_revision, expected_guard=None, published_callback=None):
    """The ordinary publisher shared by UI and a separately guarded operator.

    Provider/network I/O is outside the transaction. Internal guard runs before
    I/O and again on locked current rows; the optional trusted callback may save
    only local ledger metadata and must not perform provider/network I/O.
    """
    from .avito_services_workplace_setup import _binding
    connection, actor = _publication_connection(connection.pk, actor)
    binding = _binding(connection)
    feed = AvitoServicesFeed.objects.get(connection=connection)
    if not bound(feed, connection) or str(expected_revision) != str(feed.revision):
        raise PublicationConflict('Данные изменились. Обновите страницу.')
    if expected_guard is not None and expected_guard(connection, feed) is not True:
        raise PublicationConflict('Состав согласованного изменения изменился. Публикация остановлена.')
    public_origin()
    # Capture the precise retention authority, not just the draft revision.
    baseline = (deepcopy(feed.published), feed.published_revision, feed.published_photo_limit,
                feed.enabled, feed.public_key, deepcopy(feed.draft))
    xml_bytes(feed.draft, max_images=30)
    require_identities(feed, feed.draft)
    verified = verify_inventory(connection, feed.draft)
    if not entries(verified['sheets']):
        raise ValueError('Нет активных объявлений для публикации. Возвращение из архива требует отдельной проверки стоимости.')
    content = xml_bytes(verified['sheets'])
    media_result = validate_snapshot_media(connection, verified['sheets'],
                                          verified_observations=verified['observations'])
    with transaction.atomic():
        connection, actor = _publication_connection(connection.pk, actor, locked=True)
        locked = AvitoServicesFeed.objects.select_for_update().get(pk=feed.pk)
        current_baseline = (locked.published, locked.published_revision, locked.published_photo_limit,
                            locked.enabled, locked.public_key, locked.draft)
        if (not bound(locked, connection) or _binding(connection) != binding or locked.revision != feed.revision
                or current_baseline != baseline or _older_observation(locked, verified['observations'])):
            raise PublicationConflict('Данные изменились. Обновите страницу.')
        if expected_guard is not None and expected_guard(connection, locked) is not True:
            raise PublicationConflict('Состав согласованного изменения изменился. Публикация остановлена.')
        require_identities(locked, locked.draft)
        ensure_baseline(locked)
        before = deepcopy(locked.published)
        for item in entries(verified['sheets']):
            validate_photo_scope(connection, item['ImageUrls'])
        locked.published = verified['sheets']
        locked.published_revision = locked.revision
        locked.published_photo_limit = 10
        locked.enabled = True
        verified['observations'] = annotate_reactivation(locked.latest_observations, verified['observations'])
        locked.latest_observations = verified['observations']
        locked.last_checked_at = timezone.now()
        locked.updated_by = actor
        locked.save()
        mark_snapshot_media_published(connection, locked.published)
        retained = len(media_result.get('retained', [])) if isinstance(media_result, dict) else 0
        reason = 'Публикация файла; применение Авито ещё не подтверждено'
        if retained:
            reason += f'; {retained} прежних ссылок сохранены по опубликованному снимку, доступность источников заново не подтверждена'
        version = record_version(locked, 'publish', before, locked.published, actor=actor,
                       reason=reason, observations=verified['observations'], content=content)
        bind_publication(locked, version)
        if published_callback is not None:
            published_callback(connection, locked, version)
    return {'feed': locked, 'version': version, 'content': content,
            'observations': verified['observations'], 'media': media_result}


@login_required
@never_cache
@require_http_methods(['POST'])
def services_feed_publish(request, connection_id):
    from .communication_avito import AvitoError
    connection = connection_for(request, connection_id)
    feed = get_object_or_404(AvitoServicesFeed, connection=connection)
    try:
        publish_services_snapshot(connection, request.user, expected_revision=request.POST.get('revision'))
    except PublicationConflict as error:
        return HttpResponse(str(error), status=409)
    except (ValueError, AvitoError) as error:
        return render(request, 'pool_service/avito/services_feed.html',
                      _page_context(request, connection, feed, _safe_failure(error)), status=400)
    return redirect('avito_services_feed', connection_id=connection_id)


def _verify_publication(request, connection, feed):
    """Explicit read-only provider check; it proves status, never text or file ingestion."""
    from .communication_avito import AvitoError
    if request.POST.get('revision') != str(feed.revision):
        return HttpResponse('Данные изменились. Обновите страницу.', status=409)
    try:
        observed = verify_inventory(connection, feed.draft)['observations']
    except (ValueError, AvitoError) as error:
        return render(request, 'pool_service/avito/services_feed.html',
                      _page_context(request, connection, feed, _safe_failure(error)), status=400)
    with transaction.atomic():
        ChannelConnection.objects.select_for_update().get(pk=connection.pk)
        connection = connection_for(request, connection.pk)
        locked = AvitoServicesFeed.objects.select_for_update().get(pk=feed.pk)
        if (not bound(locked, connection) or locked.revision != feed.revision
                or _older_observation(locked, observed)):
            return HttpResponse('Данные изменились. Обновите страницу.', status=409)
        ensure_baseline(locked)
        observed = annotate_reactivation(locked.latest_observations, observed)
        locked.latest_observations = observed
        locked.last_checked_at = timezone.now()
        locked.save(update_fields=['latest_observations', 'last_checked_at'])
        record_version(locked, 'verify', locked.draft, locked.draft, actor=request.user,
                       reason='Проверка статусов и известных сроков; содержание и загрузка файла не подтверждены',
                       observations=observed)
    return redirect('avito_services_feed', connection_id=connection.pk)


@never_cache
@require_GET
def services_feed_public(request, public_key, output_format="xml"):
    feed = get_object_or_404(AvitoServicesFeed.objects.select_related('connection__channel'), public_key=public_key, enabled=True)
    connection = feed.connection
    if not feed.published or not bound(feed, connection) or not connection.is_active or not connection.channel.is_active or connection.channel.kind != 'avito':
        raise Http404
    try:
        # Only pre-upgrade snapshots can have the persisted legacy limit30.
        content = workbook_bytes(feed.published) if output_format == 'xlsx' else xml_bytes(feed.published, max_images=feed.published_photo_limit)
    except ValueError:
        return HttpResponse('Источник временно недоступен: требуется проверка данных.', status=503)
    response = HttpResponse(content, content_type=MIME if output_format == 'xlsx' else 'application/xml; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="avito-services.{output_format}"'
    response['X-Robots-Tag'] = 'noindex, nofollow, noarchive'
    return response


@never_cache
@require_GET
def services_feed_version_public(request, public_key, version_id):
    publication = get_object_or_404(AvitoServicesPublication.objects.select_related(
        'version__feed__connection__channel'), version_id=version_id, version__feed__public_key=public_key)
    version = publication.version
    feed, connection = version.feed, version.feed.connection
    if not _binding_valid(publication, connection):
        raise Http404
    try:
        content = xml_bytes(version.snapshot, max_images=10)
    except ValueError:
        return HttpResponse('Источник версии временно недоступен.', status=503)
    if hashlib.sha256(content).hexdigest() != publication.file_sha256:
        return HttpResponse('Контрольная сумма версии не совпадает.', status=503)
    response = HttpResponse(content, content_type='application/xml; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="avito-services-version-{version.pk}.xml"'
    response['X-Robots-Tag'] = 'noindex, nofollow, noarchive'
    return response
