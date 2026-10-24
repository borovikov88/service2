from io import BytesIO
from zipfile import ZipFile
import re
from xml.etree import ElementTree as ET
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, Client
from django.urls import reverse
from pool_service.models import Organization, OrganizationAccess
from pool_service.communication_models import ChannelConnection, CommunicationChannel
from pool_service.avito_feed_models import AvitoServicesFeed
from pool_service.avito_services_xml import xml_bytes
from pool_service.avito_services_feed import parse_export, workbook_bytes, verify_inventory, images


def export_bytes(ids=('111',)):
    book = Workbook(); sheet = book.active; sheet.title = 'Услуги'
    sheet.append(['Услуги'])
    sheet.append(['Id', 'AvitoId', 'Category', 'AvitoStatus', 'Title', 'Description', 'ImageUrls', 'PriceList', 'ListingFee'])
    sheet.append(['Обязательный']); sheet.append(['Подробнее'])
    for item in ids:
        sheet.append([item, item, 'Предложение услуг', 'Активно', 'Тест услуги', 'Описание', 'https://example.org/image.jpg', 'Своя услуга|Тест|100|Да|за услугу', 'Package'])
    output = BytesIO(); book.save(output); return output.getvalue()


class ServicesFeedTests(TestCase):
    def setUp(self):
        availability = patch('pool_service.avito_services_feed.validate_snapshot_media')
        availability.start()
        self.addCleanup(availability.stop)
        org = Organization.objects.create(name='Feed test')
        self.owner = User.objects.create_user('feed-owner')
        self.manager = User.objects.create_user('feed-manager')
        OrganizationAccess.objects.create(organization=org, user=self.owner, role='owner')
        OrganizationAccess.objects.create(organization=org, user=self.manager, role='manager')
        channel = CommunicationChannel.objects.create(organization=org, kind='avito', name='Avito')
        self.connection = ChannelConnection.objects.create(channel=channel, name='Services', external_id='123')
        self.url = reverse('avito_services_feed', args=[self.connection.pk])
        self.client.force_login(self.owner)

    def upload(self, **extra):
        return self.client.post(self.url, {'action': 'import', 'file': SimpleUploadedFile('123_export.xlsx', export_bytes()), **extra})

    def test_roundtrip_preserves_ids_category_fields_and_price_list(self):
        data = parse_export(export_bytes())
        self.assertEqual(parse_export(workbook_bytes(data)), data)
        with self.assertRaises(ValueError):
            parse_export(export_bytes(('111', '111')))
        data[0]['rows'][0]['Description'] = '=HYPERLINK("https://example.org")'
        book = load_workbook(BytesIO(workbook_bytes(data)), data_only=False)
        self.assertEqual(book.active['F5'].data_type, 's')

    def test_import_edit_preserves_published_snapshot_and_rejects_stale_edit(self):
        self.assertEqual(self.upload().status_code, 302)
        feed = AvitoServicesFeed.objects.get()
        feed.published = feed.draft; feed.enabled = True; feed.save()
        edit = reverse('avito_services_feed_edit', args=[self.connection.pk, '111'])
        payload = {'title': 'Новый заголовок', 'description': 'Новый текст', 'image_urls': 'https://example.org/new.jpg', 'revision': feed.revision}
        self.assertEqual(self.client.post(edit, payload).status_code, 302)
        feed.refresh_from_db()
        self.assertEqual(feed.draft[0]['rows'][0]['Title'], 'Новый заголовок')
        self.assertEqual(feed.published[0]['rows'][0]['Title'], 'Тест услуги')
        self.assertEqual(self.client.post(edit, payload).status_code, 409)
        self.assertEqual(self.client.post(edit, payload).status_code, 409)
        payload['title'] = ''
        self.assertEqual(self.client.post(edit, payload).status_code, 409)
        public = reverse('avito_services_feed_public', args=[feed.public_key])
        response = Client().get(public)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(ET.fromstring(response.content).findtext('Ad/Title'), 'Тест услуги')
        self.assertEqual(response['Content-Type'], 'application/xml; charset=utf-8')
        legacy = reverse('avito_services_feed_public_xlsx', args=[feed.public_key])
        self.assertEqual(parse_export(Client().get(legacy).content)[0]['rows'][0]['Title'], 'Тест услуги')
        self.connection.external_id = '456'; self.connection.save()
        self.assertEqual(Client().get(public).status_code, 404)
        self.assertEqual(self.client.get(self.url).status_code, 409)

    def test_bad_dimensions_and_malformed_xml(self):
        raw = export_bytes()
        def modified(replacement):
            output = BytesIO()
            with ZipFile(BytesIO(raw)) as source, ZipFile(output, 'w') as target:
                for name in source.namelist():
                    value = source.read(name)
                    if name == 'xl/worksheets/sheet1.xml':
                        value = replacement(value)
                    target.writestr(name, value)
            return output.getvalue()
        wrong = modified(lambda value: re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', value))
        self.assertEqual(parse_export(wrong), parse_export(raw))
        with self.assertRaises(ValueError): parse_export(modified(lambda value: b'<worksheet><broken'))

    def test_fresh_export_can_reconcile_active_inventory(self):
        self.upload(); feed = AvitoServicesFeed.objects.get()
        changed = SimpleUploadedFile('123_export.xlsx', export_bytes(('111', '222')))
        response = self.client.post(self.url, {'action': 'import', 'revision': feed.revision, 'file': changed})
        self.assertEqual(response.status_code, 302)
        feed.refresh_from_db()
        self.assertEqual(len(feed.draft[0]['rows']), 2)
        self.assertFalse(feed.enabled)

    def test_normalized_photo_list_limit(self):
        url = 'https://example.org/' + 'a' * 1900
        with self.assertRaises(ValueError): images('\n'.join([url] * 18))

    def test_cell_limit_rejects_truncation(self):
        data = parse_export(export_bytes())
        data[0]['rows'][0]['Description'] = 'x' * 32768
        with self.assertRaises(ValueError): workbook_bytes(data)

    def test_access_isolation_csrf_and_disabled_feed(self):
        self.client.force_login(self.manager)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.assertEqual(self.upload().status_code, 403)
        self.client.force_login(self.owner)
        self.upload()
        feed = AvitoServicesFeed.objects.get()
        self.assertEqual(Client().get(reverse('avito_services_feed_public', args=[feed.public_key])).status_code, 404)
        other = Organization.objects.create(name='Other')
        channel = CommunicationChannel.objects.create(organization=other, kind='avito', name='Other')
        foreign = ChannelConnection.objects.create(channel=channel, name='Foreign', external_id='789')
        self.assertEqual(self.client.get(reverse('avito_services_feed', args=[foreign.pk])).status_code, 404)
        csrf = Client(enforce_csrf_checks=True); csrf.force_login(self.owner)
        self.assertEqual(csrf.post(self.url, {'action': 'import'}).status_code, 403)

    @patch('pool_service.avito_services_feed.verify_inventory')
    def test_publish_is_explicit_and_uses_reviewed_revision(self, verify):
        self.upload(); feed = AvitoServicesFeed.objects.get()
        verify.return_value = {'sheets': feed.draft, 'observations': []}
        url = reverse('avito_services_feed_publish', args=[self.connection.pk])
        self.assertEqual(self.client.get(url).status_code, 405)
        self.assertEqual(self.client.post(url, {'revision': 0}).status_code, 409)
        verify.assert_not_called()
        self.assertEqual(self.client.post(url, {'revision': feed.revision}).status_code, 302)
        feed.refresh_from_db(); self.assertTrue(feed.enabled); self.assertEqual(feed.published, feed.draft)
        self.assertEqual(self.client.post(self.url, {'action': 'disable', 'revision': feed.revision}).status_code, 302)
        feed.refresh_from_db(); self.assertFalse(feed.enabled)

    @patch('pool_service.avito_workspace.fetch_item_detail', return_value={'status': 'active', 'autoload_item_id': ''})
    @patch('pool_service.avito_workspace._get', return_value={'resources': [{'id': 111}]})
    @patch('pool_service.communication_avito.access_token', return_value='test')
    @patch('pool_service.communication_avito.authorized_account_id', return_value='123')
    def test_inventory_must_match_exactly_and_preserve_autoload_id(self, account, token, get, detail):
        data = parse_export(export_bytes())
        verify_inventory(self.connection, data)
        get.assert_called_once_with('test', '/core/v1/items?per_page=100&page=1&status=active')
        get.return_value = {'resources': [{'id': 111}, {'id': 222}]}
        with self.assertRaises(ValueError): verify_inventory(self.connection, data)
        get.return_value = {'resources': [{'id': 111}]}
        detail.return_value = {'status': 'active', 'autoload_item_id': 'different'}
        with self.assertRaises(ValueError): verify_inventory(self.connection, data)

    @patch('pool_service.avito_services_feed.verify_inventory')
    def test_provider_failure_reports_only_safe_code(self, verify):
        from pool_service.communication_avito import AvitoError
        self.upload(); feed = AvitoServicesFeed.objects.get()
        url = reverse('avito_services_feed_publish', args=[self.connection.pk])
        verify.side_effect = AvitoError('provider_http_404')
        response = self.client.post(url, {'revision': feed.revision})
        self.assertContains(response, 'provider_http_404', status_code=400)
        verify.side_effect = AvitoError('secret payload not for display')
        response = self.client.post(url, {'revision': feed.revision})
        self.assertContains(response, 'provider_check_failed', status_code=400)
        self.assertNotContains(response, 'secret payload', status_code=400)
        feed.refresh_from_db()
        self.assertFalse(feed.enabled)
        self.assertEqual(feed.published, [])

    def test_xml_nested_fields_escape_content_and_omit_export_metadata(self):
        data = parse_export(export_bytes())
        data[0]['rows'][0].update(Payment='Полная предоплата | Поэтапная оплата',
            WorkDays='пн. | вт.', Title='Вода & <монтаж>',
            Description='<p>Текст</p> & детали', ImageUrls='https://example.org/a?a=1&b=2',
            CompanyName='Export only', EMail='export@example.invalid', AvitoDateEnd='2026-01-01')
        root = ET.fromstring(xml_bytes(data)); ad = root.find('Ad')
        self.assertEqual(root.attrib, {'formatVersion': '3', 'target': 'Avito.ru'})
        self.assertEqual(ad.findtext('Id'), '111'); self.assertEqual(ad.findtext('AvitoId'), '111')
        self.assertEqual(ad.findtext('Title'), 'Вода & <монтаж>')
        self.assertEqual(ad.findtext('Description'), '<p>Текст</p> & детали')
        self.assertEqual(ad.find('Images/Image').attrib['url'], 'https://example.org/a?a=1&b=2')
        self.assertEqual([o.text for o in ad.findall('Payment/Option')], ['Полная предоплата', 'Поэтапная оплата'])
        self.assertEqual([o.text for o in ad.findall('WorkDays/Option')], ['пн.', 'вт.'])
        service = ad.find('PriceList/Service')
        self.assertEqual({c.tag: c.text for c in service}, {'ServiceName': 'Своя услуга',
            'ServiceTitle': 'Тест', 'ServicePrice': '100', 'ServiceStartingPrice': 'Да', 'ServicePriceType': 'за услугу'})
        for tag in ('AvitoStatus', 'AvitoDateEnd', 'CompanyName', 'EMail', 'ImageUrls'):
            self.assertIsNone(ad.find(tag))

    def test_xml_rejects_unknown_fields_and_bad_price_rows_before_publication(self):
        for change in ({'Promo': 'XL'}, {'PriceList': 'Своя услуга||100|Да|за услугу'},
                       {'PriceList': 'Услуга||от 100||за услугу'}, {'Description': 'bad\x01text'}):
            with self.subTest(change=change):
                data = parse_export(export_bytes()); data[0]['rows'][0].update(change)
                with self.assertRaises(ValueError): xml_bytes(data)
        self.upload(); feed = AvitoServicesFeed.objects.get()
        feed.draft[0]['rows'][0]['PriceList'] = 'broken'; feed.save()
        with patch('pool_service.avito_services_feed.verify_inventory') as verify:
            result = self.client.post(reverse('avito_services_feed_publish', args=[self.connection.pk]),
                                     {'revision': feed.revision})
        self.assertEqual(result.status_code, 400); verify.assert_not_called()
        feed.refresh_from_db(); self.assertFalse(feed.enabled)
